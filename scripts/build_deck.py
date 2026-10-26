"""Build the lecture deck deck/rom-session1.pptx from docs/template.pptx.

Usage:
    python scripts/build_deck.py          # build the pptx
    python scripts/build_deck.py --pdf    # also export deck/rom-session1.pdf with PowerPoint (Windows)

Inputs:
    docs/template.pptx                     layouts and theme
    figures/anim/*.gif, *.png              animations (scripts/make_animations.py)
    figures/static/*.png                   notebook figures
    figures/thesis/T*.png                  thesis crops (scripts/extract_thesis_figs.py)
    figures/papers/*.png                   figures of an open-access paper (figures/papers/README.md)

Steps that PowerPoint cannot get from python-pptx are written in the speaker notes
as lines that start with 【アニメーション】.
"""
import argparse
import copy
import math
import re
import sys
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

TEMPLATE = ROOT / "docs" / "template.pptx"
OUT = ROOT / "deck" / "rom-session1.pptx"
ANIM = ROOT / "figures" / "anim"
STATIC = ROOT / "figures" / "static"
THESIS = ROOT / "figures" / "thesis"
PAPERS = ROOT / "figures" / "papers"
BOOK = ROOT / "figures" / "book"
BUILD = ROOT / "deck" / "_build"       # images derived for the slides only (git-ignored)

# Theme colours (spec 8.2)
DEEP = RGBColor(0x00, 0x00, 0x28)
CARD = RGBColor(0xFF, 0xFF, 0xFF)       # cards are white with an accent line (printed handout)
HEAD_FILL = RGBColor(0xD5, 0xEC, 0xEC)  # light teal for table header rows
ACC1 = RGBColor(0x00, 0x99, 0x99)
ACC2 = RGBColor(0x00, 0xD7, 0xA0)
ACC3 = RGBColor(0x00, 0xBE, 0xDC)
ACC4 = RGBColor(0x00, 0x87, 0xBE)
ACC5 = RGBColor(0x00, 0x55, 0x7C)
WARM = RGBColor(0xFF, 0xB0, 0x4A)
LIGHT = RGBColor(0xF3, 0xF3, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x5A, 0x5F, 0x73)
BLACK = RGBColor(0x00, 0x00, 0x00)
EA_FONT = "Meiryo UI"

# Geometry in cm (16:9, 33.867 x 19.05)
L = 1.14             # left edge of the content
CW = 31.6            # content width
BODY_TOP = 3.93      # top of the body placeholders
BODY_BOTTOM = 17.75  # lowest edge for content; the citation line is below
GAP = 0.3

# Layout indices (spec 8.4)
LY_TITLE, LY_CHAPTER, LY_AGENDA, LY_FREE = 37, 33, 36, 37
LY_ONE, LY_TWO, LY_THREE, LY_FOUR = 42, 47, 49, 51
LY_TABLE, LY_PROCESS, LY_STATEMENT = 55, 56, 60
DARK_LAYOUTS = set()   # the deck is printed: only light layouts

# Text colours of the current slide; new_slide() switches them for dark layouts
FG, SUBTLE = DEEP, GREY
DARK_SUBTLE = RGBColor(0xB8, 0xBC, 0xC8)

# Slide order. Keys are the slide numbers of the spec; the value is the page number shown on the slide.
ORDER = ["1", "chA", "2", "3", "4", "5", "chB", "6", "7", "8", "9", "10", "chC", "11", "12",
         "chD", "13", "14", "14b", "chE", "15", "16", "17", "18", "19", "20", "chF", "21", "22",
         "chG", "23", "24", "25a", "25b", "25c", "chH", "26", "apxA1", "apxA2", "apxA3", "apxB"]
PAGE = {k: i + 1 for i, k in enumerate(ORDER)}


def pg(*keys):
    """Page reference as shown on the slides, for example pg("5") -> "p.6"."""
    return "p." + "・".join(str(PAGE[k]) for k in keys)


# ---------------------------------------------------------------- text helpers
MARKUP = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


def plain(text):
    """Text without the markup, for width estimates."""
    return re.sub(r"\*\*(.+?)\*\*|_\{(.+?)\}|\^\{(.+?)\}", lambda m: "".join(g or "" for g in m.groups()), text)


def text_width_cm(text, pt):
    em = pt * 2.54 / 72
    return sum(em * (1.0 if ord(ch) > 0x2E80 else 0.55) for ch in plain(text))


def lines_needed(text, pt, width_cm):
    return sum(max(1, math.ceil(text_width_cm(part, pt) / (width_cm - 0.3))) for part in text.split("\n"))


def add_runs(p, text, size, color, bold=False, hl=ACC1):
    """Add runs to paragraph p. Markup: **x** bold highlight, _{x} subscript, ^{x} superscript.

    **x** can hold sub- and superscripts.
    """
    for part in MARKUP.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_runs(p, part[2:-2], size, hl, True, hl)
            continue
        r = p.add_run()
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        if part.startswith("_{"):
            r.text = part[2:-1]
            r._r.get_or_add_rPr().set("baseline", "-25000")
        elif part.startswith("^{"):
            r.text = part[2:-1]
            r._r.get_or_add_rPr().set("baseline", "30000")
        else:
            r.text = part


def set_bullet(p, color=ACC1, indent_cm=0.6):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Cm(indent_cm)))
    pPr.set("indent", str(-Cm(indent_cm)))
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def fill_tf(tf, paras, size=20, color=DEEP, bold=False, align=PP_ALIGN.LEFT, hl=ACC1,
            space_after=4, line_spacing=1.05):
    """Write paragraphs into a text frame.

    Each item is a string or a dict with keys t, size, color, bold, bullet, align, space_before.
    """
    tf.word_wrap = True
    first = True
    for item in paras:
        d = {"t": item} if isinstance(item, str) else dict(item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = d.get("align", align)
        p.line_spacing = line_spacing
        p.space_after = Pt(d.get("space_after", space_after))
        if "space_before" in d:
            p.space_before = Pt(d["space_before"])
        if d.get("bullet"):
            set_bullet(p)
        add_runs(p, d["t"], d.get("size", size), d.get("color", color), d.get("bold", bold), d.get("hl", hl))


def no_style(shape):
    """Drop the default p:style of a new autoshape, so no theme shadow or outline applies."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def textbox(slide, x, y, w, h, paras, size=20, color=None, anchor=MSO_ANCHOR.TOP, margin=0.1, **kw):
    color = FG if color is None else color
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Cm(margin))
    fill_tf(tf, paras, size=size, color=color, **kw)
    return tb


def box(slide, x, y, w, h, paras=(), fill=CARD, line=None, size=14, color=DEEP,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
        radius_cm=0.25, margin=0.2, dash=False, line_w=1.5, **kw):
    sh = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    no_style(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius_cm / min(w, h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
        if dash:
            ln = sh.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    tf = sh.text_frame
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Cm(margin))
    if paras:
        fill_tf(tf, paras, size=size, color=color, align=align, **kw)
    return sh


def card(slide, x, y, w, h, head, body, head_size=15, size=14, bullet=False, anchor=MSO_ANCHOR.TOP):
    """Card: white rounded rectangle with an accent1 line, heading 15 pt bold accent1, body 14 pt deep blue."""
    paras = []
    if head:
        paras.append({"t": head, "size": head_size, "bold": True, "color": ACC1, "space_after": 6})
    for b in body:
        d = {"t": b} if isinstance(b, str) else dict(b)
        d.setdefault("bullet", bullet)
        paras.append(d)
    return box(slide, x, y, w, h, paras, size=size, anchor=anchor, align=PP_ALIGN.LEFT,
               margin=0.35, hl=ACC1, line=ACC1, line_w=1.0)


def arrow(slide, x1, y1, x2, y2, color=ACC4, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def line(slide, x1, y1, x2, y2, color=GREY, width=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


# ---------------------------------------------------------------- slide helpers
def image_size(path):
    with Image.open(path) as im:
        return im.size


def picture(slide, path, x, y, w=None, h=None, fit=None, center=False):
    """Add a picture at (x, y). Give w or h, or fit=(w, h) to fit inside a box. Returns (w, h) in cm."""
    iw, ih = image_size(path)
    if fit:
        s = min(fit[0] / iw, fit[1] / ih)
        w, h = iw * s, ih * s
        if center:
            x += (fit[0] - w) / 2
    elif w is not None:
        h = w * ih / iw
    else:
        w = h * iw / ih
    slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))
    return w, h


def remove_placeholders(slide, keep=()):
    for ph in list(slide.placeholders):
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.TITLE or ph.placeholder_format.idx in keep:
            continue
        ph._element.getparent().remove(ph._element)


def add_slide_number(slide):
    """python-pptx does not copy the slide-number placeholder; copy it from the layout (spec 8.1)."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))
            return


def set_title(slide, head, sub=None, sub_size=16):
    """Two-line title: line 1 bold accent (layout default), line 2 regular (spec 8.3)."""
    tf = slide.shapes.title.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = head
    if sub:
        p.add_line_break()
        r2 = p.add_run()
        r2.text = sub
        r2.font.bold = False
        r2.font.size = Pt(sub_size)


def claim(slide, text, y=BODY_TOP, x=L, w=CW, size=20):
    """The one-sentence claim at the top of the body (spec 8.3). Returns the y below it."""
    n = lines_needed(text, size, w)
    h = n * size * 1.25 * 2.54 / 72 + 0.25
    textbox(slide, x, y, w, h, [{"t": text, "bold": True}], size=size, color=FG)
    return y + h + 0.25


LINKS = []   # (shape, ORDER key): citation lines that jump to an appendix page; set in build()


def cite(slide, text, to="apxA1"):
    """Short citation line at the bottom, 8 pt black. A click in the slide show opens the appendix page."""
    tb = textbox(slide, 3.2, 17.95, CW - 2.06, 0.5, [f"{text}　→ 付録 p.{PAGE[to]}"], size=8,
                 color=BLACK, anchor=MSO_ANCHOR.MIDDLE, space_after=0, line_spacing=1.0, margin=0.05)
    LINKS.append((tb, to))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs, layout, head, sub=None, number=True, keep=()):
    global FG, SUBTLE
    FG, SUBTLE = (WHITE, DARK_SUBTLE) if layout in DARK_LAYOUTS else (DEEP, GREY)
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    remove_placeholders(s, keep=keep)
    if head is not None:
        set_title(s, head, sub)
    if number:
        add_slide_number(s)
    return s


def _gif(path, x, y, max_w, max_h):
    """Size of a 16:9 GIF that fits in the box (spec 8.5 widths are upper limits). Returns (w, h)."""
    w = min(max_w, max_h * 16 / 9)
    return w, w * 9 / 16


def table(slide, x, y, col_w, rows, size=14, header_size=None, row_h=0.8, first_col_bold=True,
          header_fill=HEAD_FILL, band=(WHITE, LIGHT), placeholder=None, align_first=PP_ALIGN.LEFT, header_h=None):
    """Table with explicit colours. rows[0] is the header row."""
    nr, nc = len(rows), len(rows[0])
    if placeholder is not None:
        gf = placeholder.insert_table(nr, nc)
        gf.left, gf.top = Cm(x), Cm(y)
    else:
        gf = slide.shapes.add_table(nr, nc, Cm(x), Cm(y), Cm(sum(col_w)), Cm(row_h * nr))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Cm(w)
    for i in range(nr):
        tbl.rows[i].height = Cm(header_h if (i == 0 and header_h) else row_h)
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Cm(0.15)
            cell.margin_top = cell.margin_bottom = Cm(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
                color, bold, sz = DEEP, True, header_size or size
            else:
                cell.fill.fore_color.rgb = band[(i - 1) % 2]
                color, bold, sz = DEEP, (j == 0 and first_col_bold), size
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            parts = rows[i][j].split("\n") if rows[i][j] else [""]
            paras = [{"t": t} for t in parts]
            fill_tf(tf, paras, size=sz, color=color, bold=bold, space_after=0, line_spacing=1.0,
                    align=align_first if j == 0 else PP_ALIGN.LEFT)
    return gf


# ---------------------------------------------------------------- derived images
def make_title_image(path):
    """Vorticity of data set 1 at t = 0, cropped to the 3.48:1 picture placeholder of layout 24."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from rom_utils import CMAP_DIV, DEEP_BLUE

    d = np.load(ROOT / "data" / "cylinder_vort_f32.npz")
    x, y, w = d["x"], d["y"], d["vort"][:, :, 0]
    keep = np.abs(y) <= 1.29
    w = w[keep]
    yk = y[keep]
    vlim = np.percentile(np.abs(w), 98)
    fig = plt.figure(figsize=(12, 12 / 3.48), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.imshow(w, extent=[x[0], x[-1], yk[0], yk[-1]], origin="lower", cmap=CMAP_DIV,
              vmin=-vlim, vmax=vlim, aspect="auto", interpolation="bicubic")
    ax.add_patch(plt.Circle((0, 0), 0.5, color=DEEP_BLUE))
    ax.set_axis_off()
    fig.savefig(path)
    plt.close(fig)


def crop(src, dst, box_frac):
    """Crop a fraction box (x0, y0, x1, y1) of an image."""
    with Image.open(src) as im:
        W, H = im.size
        x0, y0, x1, y1 = box_frac
        im.crop((round(x0 * W), round(y0 * H), round(x1 * W), round(y1 * H))).save(dst)


def prepare_images():
    """Make the derived images. Thesis crops and the slide-only images stay out of git."""
    make_title_image(STATIC / "title_vorticity.png")
    crops = {
        "T1b_top.png": ("T1b.png", (0.0, 0.0, 1.0, 0.395)),      # Fig. 4.14 (a)-(d): POD modes 1-4
        "T7_t1.png": ("T7.png", (0.0, 0.0, 0.4937, 0.2015)),     # Fig. 5.20, first panel (t1, single helix)
        "T7_rows3.png": ("T7.png", (0.0, 0.0, 0.4937, 0.6002)),  # Fig. 5.20, left column t1 to t1+2dt
    }
    for dst, (src, frac) in crops.items():
        crop(THESIS / src, THESIS / dst, frac)
    BUILD.mkdir(parents=True, exist_ok=True)
    crop(PAPERS / "zhang2026_fig2.png", BUILD / "zhang2026_fig2_ring.png", (0.12, 0.15, 0.68, 0.66))
    crop_gif(ANIM / "19_manifold_vs_subspace.gif", BUILD / "19_manifold_crop.gif", (150, 0, 810, 540))


def crop_gif(src, dst, box):
    """Crop every frame of a GIF to the pixel box (x0, y0, x1, y1); keep the frame times and the loop."""
    with Image.open(src) as im:
        frames, durs = [], []
        for k in range(im.n_frames):
            im.seek(k)
            frames.append(im.convert("RGB").crop(box).quantize(colors=128, method=Image.Quantize.MEDIANCUT,
                                                                dither=Image.Dither.NONE))
            durs.append(im.info.get("duration", 80))
    frames[0].save(dst, save_all=True, append_images=frames[1:], duration=durs, loop=0, optimize=False)


# ---------------------------------------------------------------- the slides
def chapter(prs, letter, name, sub):
    s = new_slide(prs, LY_CHAPTER, None, number=False, keep=(1,))
    tf = s.shapes.title.text_frame
    p = tf.paragraphs[0]
    add_runs(p, f"{letter}  {name}", 60 if len(name) > 8 else 80, ACC1, bold=True)
    sub_ph = [ph for ph in s.placeholders if ph.placeholder_format.idx == 1][0]
    fill_tf(sub_ph.text_frame, [sub], size=24, color=DEEP)
    return s


def s01_title(prs):
    s = new_slide(prs, LY_TITLE, None, number=False)
    picture(s, STATIC / "title_vorticity.png", 0, 0, w=33.87)
    title = s.shapes.title
    title.left, title.top, title.width, title.height = Cm(L), Cm(10.3), Cm(25.6), Cm(3.4)
    tf = title.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    add_runs(p, "Reduced Order Modeling（ROM）", 40, DEEP, bold=True)
    p.add_line_break()
    add_runs(p, "流れ場の低次元化と予測", 28, ACC1)
    textbox(s, L, 14.2, 25.6, 2.6, ["流体機械演習B 第4回（3回入門シリーズの1）",
                                    "劉 志豪｜シーメンス株式会社｜2026年10月27日"], size=20)
    # Textbook cover image with its reference below it
    ch = 6.4
    cw_ = ch * image_size(BOOK / "bk2_cover.png")[0] / image_size(BOOK / "bk2_cover.png")[1]
    picture(s, BOOK / "bk2_cover.png", L + CW - cw_, 10.1, h=ch)
    textbox(s, L + CW - 6.0, 10.1 + ch + 0.05, 6.0, 1.4, [
        {"t": "教科書（以下 B&K）", "size": 11, "color": ACC1, "bold": True, "align": PP_ALIGN.RIGHT, "space_after": 0},
        {"t": "Brunton & Kutz, 2nd ed.", "size": 10, "align": PP_ALIGN.RIGHT, "space_after": 0},
        {"t": "Cambridge University Press, 2022", "size": 10, "align": PP_ALIGN.RIGHT},
    ], size=10, space_after=0, margin=0.02)
    notes(s, "表紙。教科書は1冊だけ（B&K）。上の帯は演習①で使う円柱後流（Re = 100）の渦度場。")


def s02_goals(prs):
    s = new_slide(prs, LY_AGENDA, "今日のゴール", "授業のあと、自分の研究で ROM を検討できる", keep=(12,))
    ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
    ph.left, ph.top, ph.width, ph.height = Cm(L), Cm(BODY_TOP), Cm(19.6), Cm(13.2)
    fill_tf(ph.text_frame, [
        {"t": "授業後に「自分の研究で ROM を検討できる」状態になる", "bold": True, "space_after": 18},
        {"t": "① どういう時に ROM を考えるか、判断できる", "space_after": 12},
        {"t": "② メリット・デメリットを言える", "space_after": 12},
        {"t": "③ どこまで期待できるか、見積もれる", "space_after": 12},
        {"t": "④ 実際に着手する時に、何を準備するか分かる", "space_after": 18},
        {"t": "ROM（Reduced Order Model、低次元モデル）：高忠実度の解析を、"
              "ずっと少ない自由度で近似するモデル", "size": 16, "color": GREY},
    ], size=22)
    card(s, 21.6, BODY_TOP, 11.1, 8.4, "今日の流れ（100分）", [
        "A　動機・「削減」の意味　10分",
        "B　SVD → POD / PCA　15分",
        "C　DMD　8分",
        "D　Colab 演習　15分",
        "E　手法マップ　12分",
        "F　歴史と最前線　7分",
        "G　ワークフローと思考実験　25分",
        "H　まとめ・質疑　8分",
    ], size=15)
    notes(s, f"4つのゴールを読み上げる。最後のまとめ（{pg('26')}）で、1つずつ答えを返す。\n"
             "【アニメーション】ゴール①〜④：上から1行ずつ表示")


def s03_why(prs):
    s = new_slide(prs, LY_ONE, "なぜ ROM か", "1回の解析が重すぎて、必要な回数を回せない")
    y = claim(s, "1回の高忠実度解析が重すぎて、設計で必要な回数を回せない")
    # left: 3D CFD (thesis, Fig. 5.20 first panel)
    ih = 5.9
    w, h = picture(s, THESIS / "T7_t1.png", L, y - 0.2, fit=(10.6, ih), center=True)
    textbox(s, L - 0.3, y + ih - 0.15, 11.2, 1.5, [
        {"t": "3D CFD（高忠実度）", "bold": True, "align": PP_ALIGN.CENTER, "space_after": 0},
        {"t": "例：博論の渦ロープ解析。節点 約1000万 × 5万ステップ", "size": 13, "align": PP_ALIGN.CENTER},
    ], size=16)
    my = y + 2.6   # middle line of the row
    # middle: ROM
    arrow(s, 12.0, my, 13.6, my, width=3)
    box(s, 13.8, my - 1.4, 6.4, 2.8, [
        {"t": "ROM", "size": 26, "bold": True, "space_after": 0},
        {"t": "自由度 r ≪ n の近似モデル", "size": 14},
    ], fill=HEAD_FILL, line=ACC1, line_w=1.5)
    arrow(s, 20.4, my, 22.5, my, width=3)
    # right: 1D system model as a chain of components
    comps = ["ポンプ", "管路", "水車", "制御"]
    cx = 22.7
    for k, c in enumerate(comps):
        bx = cx + k * 2.55
        box(s, bx, my - 0.7, 2.15, 1.4, [c], fill=None, line=ACC4, size=12, color=DEEP, radius_cm=0.15)
        if k < len(comps) - 1:
            line(s, bx + 2.15, my, bx + 2.55, my, color=ACC4, width=2)
    textbox(s, 22.7, my + 0.9, 10.0, 1.5, [
        {"t": "1D 系統モデル", "bold": True, "align": PP_ALIGN.CENTER, "space_after": 0},
        {"t": "部品をつないだ常微分方程式。速い", "size": 14, "align": PP_ALIGN.CENTER},
    ], size=16)
    # bottom: the three situations
    cy = y + ih + 1.5
    ch = BODY_BOTTOM - cy
    cw = (CW - 2 * GAP) / 3
    card(s, L, cy, cw, ch, "many-query：何度も回す", [
        "最適化・不確かさの定量化（UQ）で、条件を変えて同じ解析を何度も回す",
    ])
    card(s, L + cw + GAP, cy, cw, ch, "real-time：実時間で回す", [
        "Digital Twin・HIL（実物の制御器を模擬モデルとつないで試す）",
        "例：バッテリー熱の 3D 解析を ROM にして、急速充電1サイクル（2225 s）の計算が "
        "**3.92 h → 0.04 s**（Siemens, 2023）",
    ])
    card(s, L + 2 * (cw + GAP), cy, cw, ch, "co-simulation：1D と 3D をつなぐ", [
        "系統モデルに 3D の現象を部品として入れる",
        "例：EPFL は水力プラントの 1D モデル（SIMSEN）に、渦心振れ回りの励振源を注入した",
    ])
    cite(s, "Benner+ 2015｜Siemens 2023｜Alligné+ 2016｜図：博論 Fig. 5.20（p.146）、数値 p.142–143")
    notes(s, "流体機械の設計では、同じ解析を何十回も回したい場面が多い。3つの場面を左から。\n"
             "左の図は博論のドラフトチューブの CFD（ANSYS CFX、SBES）。節点 約1000万、時間刻みは振れ回り1周期の"
             "1/1000、振れ回り50周期分＝約5万ステップ（博論 p.142–143）。等圧面で渦ロープを描いたもの。\n"
             "数値（3.92 h → 0.04 s）はベンダーの自己申告（Siemens のブログ）であることを一言添える。\n"
             "【アニメーション】下の3枚のカード：左から順に")


def s04_lowrank(prs):
    s = new_slide(prs, LY_ONE, "次元を削減する", "本当の意味は「低ランク構造」")
    y = claim(s, "ROM が成立するのは、高次元の状態データが実は低ランク構造（r ≪ n）を持つから")
    top = y + 0.8
    H = 7.5
    box(s, 1.8, top, 3.6, H, [{"t": "X", "size": 28, "bold": True}], fill=ACC4, radius_cm=0.05)
    textbox(s, 1.3, top + H + 0.2, 4.6, 1.6, [{"t": "n × m", "bold": True, "align": PP_ALIGN.CENTER},
                                              {"t": "列＝1時刻の場", "size": 14, "align": PP_ALIGN.CENTER}], size=16)
    textbox(s, 5.6, top + H / 2 - 0.8, 1.2, 1.6, [{"t": "≈", "size": 36, "align": PP_ALIGN.CENTER}],
            anchor=MSO_ANCHOR.MIDDLE)
    box(s, 7.0, top, 1.0, H, [], fill=ACC1, radius_cm=0.05)
    textbox(s, 6.2, top + H + 0.2, 2.6, 1.6, [{"t": "n × r", "bold": True, "align": PP_ALIGN.CENTER},
                                              {"t": "r 本の列", "size": 14, "align": PP_ALIGN.CENTER}], size=16)
    textbox(s, 8.2, top - 0.2, 1.0, 1.4, [{"t": "×", "size": 28, "align": PP_ALIGN.CENTER}],
            anchor=MSO_ANCHOR.MIDDLE)
    box(s, 9.4, top, 3.6, 1.0, [], fill=WARM, radius_cm=0.05)
    textbox(s, 9.0, top + 1.2, 4.4, 1.6, [{"t": "r × m", "bold": True, "align": PP_ALIGN.CENTER},
                                          {"t": "r 本の行", "size": 14, "align": PP_ALIGN.CENTER}], size=16)
    textbox(s, 8.4, top + 3.0, 6.6, 3.0, [
        {"t": "列＝空間パターン（形）", "size": 14, "color": ACC1, "bold": True},
        {"t": "行＝時間の重み", "size": 14, "color": RGBColor(0xC0, 0x7A, 0x1A), "bold": True},
        {"t": "r ≪ n なら、少ない数で場全体を表せる", "size": 14},
    ], size=14)
    x0 = 15.4
    textbox(s, x0, y + 0.3, CW + L - x0, BODY_BOTTOM - y - 0.3, [
        {"t": "**低ランク仮説**：n 個の格子点の値は、独立に動いていない。少数の空間パターンの重ね合わせで書ける",
         "space_after": 16},
        {"t": "**注意**：「次元の呪い」（Bellman 1957）は別の概念。入力パラメータの次元が増えると、"
              "必要なサンプル数が指数的に増える、という話", "space_after": 16},
        {"t": "ROM では、パラメトリック ROM の学習データ数が爆発する、という**限界の側**に現れる"
              "（前回講義の「代償：データを大量に使う」と同じ話）"},
    ], size=20)
    cite(s, "B&K Ch.1｜Bellman 1957")
    notes(s, "「削減」は自由度を捨てることではなく、データが少数のパターンで書けることを使うこと。\n"
             "次元の呪いと混同しやすいので、ここで区別しておく。\n"
             "【アニメーション】右の3項目：上から順に")


def s05_quadrant(prs):
    s = new_slide(prs, LY_TABLE, "何が変わるか × 線形/非線形", "自分の問題がどの象限かで、手法が決まる", keep=(12,))
    y = claim(s, "自分の問題がどの象限かで、使う手法が決まる")
    ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
    rows = [
        ["", "時間のみ\n（固定運転点の非定常）", "パラメータのみ\n（定常のスイープ）", "両方\n（運転点をまたぐ非定常）"],
        ["主な力学の作り方", "DMD、Galerkin 射影", "POD ＋ 補間 / GP", "OpInf、LaSDI、NN"],
        ["線形\n（重ね合わせが効く）", "—", "—", "バッテリー熱\n（LTI / LPV：時間応答 ＋ 流量で補間）"],
        ["非線形\n（移流・渦が主役）", "固定運転点の PVC\n（振れ回り渦）", "—", "運転点をまたぐ PVC"],
    ]
    table(s, L, y + 0.2, [7.0, 8.2, 8.2, 8.2], rows, size=18, row_h=1.8, placeholder=ph)
    textbox(s, L, BODY_BOTTOM - 2.2, CW, 2.2, [
        {"t": "パラメータ：1回の計算の中では変えない条件（流量・回転数・形状・入口温度など）。"
              "時間：1回の計算の中で進む変数", "size": 14, "color": GREY},
        {"t": "PVC（precessing vortex core）：旋回流の渦心が軸のまわりを回る振れ回り。"
              f"GP：ガウス過程回帰。OpInf は {pg('18')}、LaSDI は {pg('20')}・{pg('22')} で説明", "size": 14, "color": GREY},
        {"t": "表は Benner et al. (2015) の分類を簡略化したもの", "size": 14, "color": GREY},
    ], size=14)
    cite(s, "Benner+ 2015（分類を簡略化）")
    notes(s, "列＝何が変わるか、行＝線形か非線形か。例題を3つ置いた。自分の研究がどこに入るか考えてもらう。\n"
             "バッテリー熱は熱伝導がほぼ線形なので LTI が使える。PVC は渦の移流なので非線形。\n"
             "【アニメーション】表の行：主な力学 → 線形 → 非線形 の順に")


def s06_svd(prs):
    s = new_slide(prs, LY_TWO, "SVD", "行列を3つに分ける")
    y = claim(s, "データを行列 X（n×m、列＝1時刻の場）に並べ、X = UΣV^{T} に分解する")
    textbox(s, L, y, 15.2, 3.6, [
        {"t": "n：空間点の数（× 変数の数）", "bullet": True},
        {"t": "m：スナップショット（ある時刻の場）の数", "bullet": True},
        {"t": "流体では n ≫ m。列＝スナップショット", "bullet": True},
    ], size=20)
    textbox(s, L, y + 3.6, 15.2, 1.0, [
        {"t": "SVD（特異値分解）：どんな行列も U・Σ・V^{T} の積に分けられる", "size": 14, "color": GREY},
    ], size=14)
    # left bottom: the thesis figure, same reshaping
    ty = y + 4.6
    tw, th = picture(s, THESIS / "T0.png", L, ty, fit=(15.2, BODY_BOTTOM - ty))
    textbox(s, L + tw + 0.3, ty, 15.2 - tw - 0.3, BODY_BOTTOM - ty, [
        {"t": "博論でも同じ整形", "bold": True, "size": 16},
        {"t": "PIV の速度場・圧力場を列に並べて X にする。多変数は縦に積む（博論 Fig. 4.1）", "size": 14},
    ], size=14, anchor=MSO_ANCHOR.MIDDLE)
    # right: GIF 06, then X = U S V^T with the shape ratio n : m = 3 : 1
    gx = 17.54
    w, h = _gif(ANIM / "06_build_X.gif", gx, y, 15.2, 8.6)
    picture(s, ANIM / "06_build_X.gif", gx, y, w=w)
    top, u = y + h + 0.3, 0.95
    xs = gx + 0.4
    items = [("X", 1, 3, ACC4, "n×m"), ("=", None, None, None, None), ("U", 1, 3, ACC1, "n×m"),
             ("Σ", 1, 1, WARM, "m×m|対角"), ("V^{T}", 1, 1, ACC5, "m×m")]
    for name, cw_, ch_, col, lab in items:
        if cw_ is None:
            textbox(s, xs, top + 1.0 * u, 0.8, 1.0, [{"t": "=", "size": 24, "align": PP_ALIGN.CENTER}],
                    anchor=MSO_ANCHOR.MIDDLE)
            xs += 0.9
            continue
        box(s, xs, top, cw_ * u, ch_ * u, [{"t": name, "size": 16, "bold": True}], fill=col, radius_cm=0.04,
            margin=0.02)
        if cw_ * u > 0 and lab:
            textbox(s, xs + cw_ * u + 0.05, top + (ch_ * u - 1.0 if ch_ > 1 else 0), 1.5, 1.0,
                    [{"t": t, "space_after": 0} for t in lab.split("|")], size=11, anchor=MSO_ANCHOR.BOTTOM)
        xs += cw_ * u + 1.45
    textbox(s, xs - 0.2, top, L + CW - xs + 0.2, 3 * u, [
        {"t": "模式図（n : m を 3 : 1 に縮めて描画）。演習データは n = 22,500、m = 151", "size": 13,
         "color": GREY},
    ], size=13, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "B&K §1.1–1.2｜博論 Fig. 4.1")
    notes(s, "GIF 06：渦度場が1時刻ずつ現れ、行列 X の列として積まれていく。「列＝スナップショット」を目で確認する。\n"
             "演習データ：100 × 225 格子 = 22,500 点、151 時刻。")


def side_text(slide, x, y, paras, size=20, w=None):
    w = w if w is not None else L + CW - x
    textbox(slide, x, y, w, BODY_BOTTOM - y, paras, size=size)


def one_gif_slide(prs, head, sub, claim_text, gif, paras, cite_text, note, gif_w=19.6, size=20):
    s = new_slide(prs, LY_ONE, head, sub)
    y = claim(s, claim_text)
    w, h = _gif(gif, L, y, gif_w, BODY_BOTTOM - y)
    picture(s, gif, L, y, w=w)
    side_text(s, L + w + 0.5, y, paras, size=size)
    cite(s, cite_text)
    notes(s, note)
    return s, y, w, h


def s07_meaning(prs):
    one_gif_slide(
        prs, "U・Σ・V の意味", "空間モード × 振幅 × 時間係数",
        "U の列＝空間モード、σ_{i}＝そのモードの大きさ、V の列＝時間変化（時間係数 a_{i} = σ_{i} v_{i}）",
        ANIM / "07_modes_and_coeffs.gif",
        [{"t": "U：空間で正規直交。形は固定", "bullet": True},
         {"t": "V：時間で正規直交。重みが動く", "bullet": True, "space_after": 14},
         {"t": "x_{k} = Σ_{i} σ_{i} u_{i} v_{ik}", "size": 24, "bold": True, "color": ACC1,
          "align": PP_ALIGN.CENTER, "space_after": 14},
         {"t": "k 時刻の場は、空間モードを時間係数で重み付けして足したもの", "space_after": 14},
         {"t": f"モード1・2 は対になり、時間係数 a_{{1}}–a_{{2}} は円を描く（{pg('10')}）", "size": 16, "color": GREY}],
        "B&K §1.1",
        "GIF 07：左の渦度場は動くが、右上の空間モードは動かない。右下の時間係数だけが動き、位相平面に円を描く。\n"
        "正規直交：互いに直交し、長さが1。\n"
        "σᵢ² をモードの「エネルギー」と呼ぶ。速度場なら運動エネルギー、演習の渦度場ならエンストロフィー"
        "（渦度の2乗の積分）に当たる。")


def s08_why(prs):
    s, y, w, h = one_gif_slide(
        prs, "なぜそうなるか", "空間側と時間側から見た、同じ相関",
        "U と V は、同じ相関構造を「空間側」と「時間側」から見たもの",
        ANIM / "08_U_from_snapshots.gif",
        [], "Sirovich 1987｜B&K §1.3",
        "GIF 08：スナップショットが時間係数の重みで順に足され、空間モード1が浮かび上がる（U = XVΣ⁻¹）。\n"
        "最初のフレームは足し終わった結果。\n"
        "【アニメーション】右の項目：上から順に",
        gif_w=18.4)
    x = L + w + 0.5
    # the two correlation matrices, drawn to scale (n > m)
    box(s, x, y, 2.8, 2.8, [{"t": "XX^{T}", "size": 16, "bold": True}, {"t": "n×n", "size": 12}], fill=ACC4,
        radius_cm=0.04)
    box(s, x + 3.3, y + 1.8, 1.0, 1.0, [], fill=WARM, radius_cm=0.04)
    textbox(s, x + 4.4, y + 1.6, 3.0, 1.3, [{"t": "X^{T}X", "size": 16, "bold": True},
                                            {"t": "m×m", "size": 12}], size=16, anchor=MSO_ANCHOR.MIDDLE)
    side_text(s, x, y + 3.2, [
        {"t": "XX^{T}（空間相関）の固有ベクトル ＝ U", "bullet": True},
        {"t": "X^{T}X（時間相関）の固有ベクトル ＝ V", "bullet": True},
        {"t": "固有値はどちらも σ_{i}^{2}", "bullet": True},
        {"t": "n ≫ m なので小さい X^{T}X を解く → **method of snapshots**（Sirovich 1987）", "bullet": True},
        {"t": "**U = XVΣ^{−1}**：空間モードは、スナップショットを時間係数で重み付けして足したもの。"
              "これが直感の核", "bullet": True},
    ], size=18)


def s09_pod_pca(prs):
    s = new_slide(prs, LY_TWO, "POD と PCA", "同じ SVD。違いは前処理と行列の向き")
    y = claim(s, "同じ SVD。違いは「平均を引くか」「スケーリングするか」「行列の向き」")
    textbox(s, L, y, 15.2, BODY_BOTTOM - y, [
        {"t": "**PCA**（主成分分析、統計）：平均を引き、必要なら分散で正規化して SVD。"
              "主成分＝右特異ベクトル", "space_after": 14},
        {"t": "**POD**（固有直交分解、流体）：平均を引くと揺らぎのモード。引かないと第1モードが平均場になる",
         "space_after": 14},
        {"t": "**行列の向きの罠**：sklearn の PCA は samples × features（行＝時刻）。"
              "転置すると U と V の意味が入れ替わる → 演習で必ず確認"},
    ], size=20)
    gx = 17.54
    # X (n x m): a column is a snapshot; U holds the spatial modes
    box(s, gx + 0.3, y + 0.2, 1.8, 5.4, [], fill=LIGHT, line=ACC4, radius_cm=0.04)
    box(s, gx + 0.3, y + 0.2, 0.45, 5.4, [], fill=ACC1, radius_cm=0.02)
    textbox(s, gx + 2.3, y + 0.2, 5.4, 5.4, [
        {"t": "X：n × m", "bold": True, "size": 18},
        {"t": "列＝1時刻の場", "size": 14},
        {"t": "空間モード＝**U の列**", "size": 16},
        {"t": "（B&K・本講）", "size": 13, "color": GREY},
    ], size=16, anchor=MSO_ANCHOR.MIDDLE)
    box(s, gx + 7.9, y + 1.9, 5.4, 1.8, [], fill=LIGHT, line=WARM, radius_cm=0.04)
    box(s, gx + 7.9, y + 1.9, 5.4, 0.45, [], fill=WARM, radius_cm=0.02)
    textbox(s, gx + 7.6, y + 3.6, 7.6, 3.0, [
        {"t": "X^{T}：m × n", "bold": True, "size": 18},
        {"t": "行＝1時刻の場（sklearn の形）", "size": 14},
        {"t": "X^{T} を SVD → 空間モード＝**V の列**", "size": 16},
        {"t": "（sklearn では components_）", "size": 13, "color": GREY},
    ], size=16)
    ny = y + 7.0
    w, h = picture(s, STATIC / "nb02_mean_subtraction.png", gx, ny, fit=(15.2, BODY_BOTTOM - ny - 0.8))
    textbox(s, gx, ny + h + 0.05, 15.2, 0.8, [{"t": "演習：平均を引く／引かないで、第1モードが変わる", "size": 14,
                                              "color": GREY}], size=14)
    cite(s, "B&K §1.5")
    notes(s, "PCA と POD は同じ計算。違うのは前処理と、行列を立てるか寝かせるか。\n"
             "sklearn の PCA に X（n×m）をそのまま渡すと、空間の点を標本とみなすので、components_ には時間方向のモードが出る。引かれる平均も時間平均の場ではなく、各時刻の空間平均になる。空間モードがほしいときは X.T を渡す。\n"
             "【アニメーション】左の3項目：上から順に")


def s10_truncation(prs):
    s = new_slide(prs, LY_TWO, "打ち切りと落とし穴", "何モード残すか。移動する構造に注意")
    y = claim(s, "何モード残すかは、エネルギー比・最良近似・硬閾値で決める。ただし移動する構造に注意")
    cw = (CW - 2 * GAP) / 3
    gw = cw
    gh = gw * 9 / 16
    cols = [
        ("10a_rank_r_recon.gif", "① 打ち切り：何モード残すか", [
            "r を増やすと、再構成の誤差が下がる",
            "SVD の打ち切り＝rank-r の最良近似（Eckart–Young）",
            "決め方：累積エネルギー 90 / 95 / 99%。ノイズ下は Gavish–Donoho 硬閾値",
        ]),
        ("10b_rotating_pattern.gif", "② 落とし穴：回る模様", [
            "cos(θ−ωt) = cosθ cosωt + sinθ sinωt",
            "→ 固定モード2本 × 時間係数2本。σ_{1} ≈ σ_{2} の**モード対**、a_{1}–a_{2} は円",
            "1運転点の PVC も、軌道半径が小さいと壁面では m = 1 が支配的 → モード対1組",
        ]),
        ("10c_translating_bump.gif", "③ 落とし穴：移動する山", [
            "形を保って動く山は、位置ごとに新しいモードが要る → σ がゆっくりしか減らない",
            f"運転点をまたぐ、または単螺旋⇄二重螺旋が切り替わる（博論 Fig. 5.20 → {pg('19')}）と、PVC でも同じ",
            "→ **Kolmogorov barrier** の入口",
        ]),
    ]
    for k, (gif, head, body) in enumerate(cols):
        x = L + k * (cw + GAP)
        picture(s, ANIM / gif, x + (cw - gw) / 2, y - 0.1, w=gw)
        cy = y + gh + 0.35
        card(s, x, cy, cw, BODY_BOTTOM - cy, head,
             [{"t": b, "bullet": True, "space_after": 3} for b in body], size=14)
    cite(s, "B&K §1.7, §12.5｜博論 Fig. 5.20")
    notes(s, "3本の GIF は同時に動く。左から順に指しながら説明する。\n"
             "① 10a：r を増やすほど誤差が下がる。SVD の打ち切りが rank-r の最良近似（Eckart–Young）。\n"
             "② 10b：回るだけの模様は、固定の形2つを cos と sin で混ぜれば書ける。だから σ が2本ずつ並ぶ。"
             "PVC が1つの運転点で回っているだけなら、これと同じでモード対1組で足りる。"
             "周方向の高調波 m を含めても、モードは 2×（有効な m の数）本で済む。\n"
             "③ 10c：形を変えずに動く山は、SVD から見ると毎時刻違う形。σ がなかなか減らない。"
             "運転点をまたいで PVC の形が変わると同じことが起きる。博論 Fig. 5.20 は、同じ条件で単螺旋と二重螺旋が"
             "周期的に切り替わる例。1つの線形基底で両方の状態を表すと、モードが倍以上要る。"
             f"これが {pg('19')} の Kolmogorov barrier につながる。\n"
             "m = 1：周方向の波数が1（壁面の圧力が1周で1回だけ高くなる形）。軌道半径が小さいときに成り立つ。")


def s11_dmd1(prs):
    one_gif_slide(
        prs, "DMD①", "データから線形の力学を取り出す",
        "X′ ≈ AX を最小二乗で解く。SVD を再利用するので、POD の続きでできる",
        ANIM / "11_dmd_rollout.gif",
        [{"t": "DMD（動的モード分解）：各モードが1つの周波数と成長率を持つ分解", "size": 16, "color": GREY,
          "space_after": 10},
         {"t": "X = [x_{1} … x_{m−1}]、X′ = [x_{2} … x_{m}]（1ステップずらす）", "bullet": True},
         {"t": "Ã = U^{T}X′VΣ^{−1}（r × r）", "bullet": True},
         {"t": "Ã の固有値 λ_{i} → 角周波数 arg(λ_{i})/Δt、成長率 ln|λ_{i}|/Δt", "bullet": True},
         {"t": "DMD モード Φ = X′VΣ^{−1}W（exact DMD、W は Ã の固有ベクトル）", "bullet": True},
         {"t": "これ自体が線形のデータ駆動 ROM：x_{k} ≈ Σ_{i} φ_{i} b_{i} λ_{i}^{k}", "bullet": True}],
        "Schmid 2010｜Tu+ 2014｜B&K §7.2",
        "GIF 11：左が真の渦度場、右が DMD モードの重ね合わせで進めた場。下の固有値は単位円の上に並ぶ（周期的な流れ）。\n"
        "最初のフレームから固有値プロットが見えている。",
        gif_w=18.4, size=18)


def s12_dmd2(prs):
    s = new_slide(prs, LY_ONE, "DMD②", "POD との違いと、ノイズの落とし穴")
    y = claim(s, "POD はエネルギー順（周波数は混ざりうる）、DMD は周波数ごと（各モードが単一周波数＋成長率）")
    gw = 11.5
    picture(s, ANIM / "12_noise_eigs.gif", L, y, w=gw)
    gy = y + gw * 9 / 16 + 0.2
    card(s, L, gy, gw, BODY_BOTTOM - gy, "落とし穴：ノイズで固有値が内側（減衰側）へ", [
        {"t": "偏り（Dawson 2016）は fbDMD・tlsDMD で補正", "bullet": True},
        {"t": "ランク r：「注目したい周波数成分が明確に特定できるぐらいまで…試行錯誤して決めてもいい」"
              "「95%と設定しても，ノイズ成分の DMD モード振幅が卓越する場合が多い」（博論 p.93）",
         "bullet": True, "size": 12, "color": GREY},
    ], size=13)
    rx = L + gw + 0.5
    rw = L + CW - rx
    th = 6.7
    w1, _ = picture(s, THESIS / "T1b_top.png", rx, y, h=th)
    picture(s, THESIS / "T2b.png", rx + w1 + 0.3, y, h=th)
    textbox(s, rx, y + th + 0.02, rw, 0.6, [
        {"t": "博論の実測：POD 係数に f_{PVC} と f_{2ndPVC} が混在（Fig. 4.14）／DMD は周波数ごと（Fig. 4.16）",
         "size": 12, "color": GREY},
    ], size=12)
    ny = y + th + 0.75
    w, h = picture(s, STATIC / "nb09_pod_fft_vs_dmd.png", rx, ny, h=BODY_BOTTOM - ny)
    textbox(s, rx + w + 0.3, ny, rw - w - 0.3, BODY_BOTTOM - ny, [
        {"t": "演習①：きれいな円柱では POD 係数もほぼ1周波数。DMD は固有値として離散的に出る", "size": 14},
        {"t": "周波数が多数ある流れでは、位相平均より POD・DMD のモード分解が妥当（博論 p.119）", "size": 13,
         "color": GREY},
    ], size=14, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "Schmid 2022｜Kutz+ 2016｜Dawson+ 2016｜博論 §4.2.2, Fig. 4.14, 4.16", to="apxA2")
    notes(s, "GIF 12：円柱データにノイズを 0 → 揺らぎ RMS の 200% まで足していく。r = 7（エネルギー 99%）。"
             "固有値が単位円の内側（減衰側）に流れる。右のパネルは |λ| の拡大。\n"
             "fbDMD：前向きと後ろ向きの DMD を組み合わせる。tlsDMD：全最小二乗で解く。どちらもノイズの偏りを減らす。\n"
             "PSD：パワースペクトル密度。")


def s13_ex1(prs):
    s = new_slide(prs, LY_TWO, "演習①", "円柱後流（解析結果の処理）")
    y = claim(s, "理想データで「手法が正しい」ことを確かめる")
    textbox(s, L, y, 21.5, 0.9, [
        {"t": "**手順（Colab）**：① clone → ② load（X を作る）→ ③ SVD → ④ モード → ⑤ DMD", "size": 16},
    ], size=16, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, L + 21.5, y, CW - 21.5, 0.9, [
        {"t": "github.com/ShigoZhihao/waseda-course-rom-1", "size": 12, "color": GREY, "align": PP_ALIGN.RIGHT},
    ], size=12, anchor=MSO_ANCHOR.MIDDLE)
    gy = y + 1.1
    cap_h = 0.8
    figs = [("nb04_sigma_energy.png", "見る：σ の急な減衰"), ("nb05_coeffs_phase.png", "見る：モード対（σ_{1} ≈ σ_{2}）と a_{1}–a_{2} の円"),
            ("nb08_dmd.png", "見る：DMD の周波数 ＝ St"), ("nb05_modes.png", "空間モード 1〜4（2つずつ対）")]
    asp = [image_size(STATIC / f)[0] / image_size(STATIC / f)[1] for f, _ in figs]
    tall = max(1 / asp[0] + 1 / asp[2], 1 / asp[1] + 1 / asp[3])
    cw2 = min((CW - 0.6) / 2, (BODY_BOTTOM - gy - 2 * cap_h - 0.2) / tall)
    gx = L + (CW - 2 * cw2 - 0.6) / 2
    yy = [gy, gy]
    for k, (f, cap) in enumerate(figs):
        col = k % 2
        x = gx + col * (cw2 + 0.6)
        w, h = picture(s, STATIC / f, x, yy[col], w=cw2)
        textbox(s, x, yy[col] + h, cw2, cap_h, [{"t": cap, "size": 14}], size=14)
        yy[col] += h + cap_h + 0.2
    cite(s, "データ：Kutz+ 2016（円柱後流 Re = 100、dmdbook.com）｜B&K §1, §7.2")
    notes(s, "ここで Colab を開いて、セル0から順に実行する（約2分）。▶ のセルはスライドのアニメーションなので飛ばしてよい。\n"
             "ストローハル数 St = fD/U：渦の放出周波数を無次元にしたもの。")


def s14_ex2(prs):
    s = new_slide(prs, LY_TWO, "演習②", "PIV 風に汚したデータで、何が壊れるか")
    y = claim(s, "同じ関数を「実験のように汚したデータ」に当てる。真値があるので、壊れ方を測れる")
    gw = 20.4
    picture(s, ANIM / "14b_gappy_fill.gif", L, y, w=gw)
    rx = L + gw + 0.4
    rw = L + CW - rx
    ch = (BODY_BOTTOM - y - GAP) / 2
    card(s, rx, y, rw, ch, "4つの劣化", [
        {"t": t, "bullet": True, "space_after": 3} for t in [
            "間引き：元の格子で4点に1点", "ノイズ：揺らぎ RMS の 5%", "欠損：5% の点 ＋ 楕円3個", "円柱の内側をマスク"]
    ], size=15)
    card(s, rx, y + ch + GAP, rw, ch, "何が起きるか", [
        {"t": t, "bullet": True, "space_after": 3} for t in [
            "欠損の埋め方で σ が変わる", "ノイズで弱い DMD モードの固有値が内側へ",
            "σ の減衰が緩くなり、何モード残すかが自明でなくなる"]
    ], size=15)
    cite(s, "Everson & Sirovich 1995｜Dawson+ 2016", to="apxA1")
    notes(s, "GIF 14b：欠損だらけの場が Gappy POD の反復で1回ずつ埋まり、真値との誤差が下がる。\n"
             "左上の黒い輪郭が欠損。右の2本のグラフは σ の並び（真値との比較）と、欠損部の誤差の下がり方。")


def s14b_thesis(prs):
    s = new_slide(prs, LY_TWO, "実験データでは", "博論の PIV 計測で起きたこと")
    y = claim(s, "本物の PIV 計測でも、σ（TKE）の減衰とモード対を見て、残すモードを決める")
    th = 9.4
    w1 = th * image_size(THESIS / "T1a.png")[0] / image_size(THESIS / "T1a.png")[1]
    w2 = th * image_size(THESIS / "T1b_top.png")[0] / image_size(THESIS / "T1b_top.png")[1]
    x0 = L + (CW - w1 - w2 - 0.4) / 2
    picture(s, THESIS / "T1a.png", x0, y, h=th)
    picture(s, THESIS / "T1b_top.png", x0 + w1 + 0.4, y, h=th)
    cy = y + th + 0.05
    textbox(s, x0, cy, w1, 0.8, [{"t": "σ（TKE）の減衰。TKE ≥ 1% を支配的と定義（Fig. 4.12）", "size": 13,
                                  "color": GREY}], size=13)
    textbox(s, x0 + w1 + 0.4, cy, w2, 0.8, [{"t": "モード1〜4：隣り合うモードが対。時間係数に複数の周波数（Fig. 4.14）",
                                             "size": 13, "color": GREY}], size=13)
    by = cy + 0.85
    bh = BODY_BOTTOM - by
    textbox(s, L, by, CW, bh, [
        {"t": "演習の合成データ（振れ回り渦と壁面の圧力）は、軌道半径 r^{*} = 0.13 を博論 Fig. 3.12 から読んで作った",
         "size": 14},
        {"t": "実験条件：PIV 1500 Hz・16377 ペア、壁面圧力 25 kHz で同期計測（博論 p.42、Fig. 3.9）", "size": 13,
         "color": GREY},
    ], size=14, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "博論 Fig. 3.9, 3.12, 4.12, 4.14, p.42, p.93", to="apxA3")
    notes(s, "演習②で体験した壊れ方は、博論のデータ処理でも起きた。ランクは 95% 基準ではなく試行錯誤で決めた（博論 p.93）。\n"
             "TKE：乱れの運動エネルギー。左の図で、TKE が 1% 以上のモードを支配的とした。\n"
             "右の図：隣り合うモードが対になる（モード1・2、3・4）。時間係数のスペクトルには複数のピークがある。\n"
             "14a の動く版は figures/anim/14a_pvc_orbit.gif。")


def s15_basis_dynamics(prs):
    s = new_slide(prs, LY_THREE, "ROM ＝ 基底 ＋ 力学", "以降の手法は、すべてこの組み合わせ", keep=(1, 2, 12))
    y = claim(s, "以降の手法はすべて「基底をどう選ぶか」「力学をどう作るか」の組み合わせ")
    phs = {p.placeholder_format.idx: p for p in s.placeholders}
    content = [
        (1, "基底", "どの低次元座標で表すか", ["線形：POD / SVD", "非線形：autoencoder、二次多様体"]),
        (2, "力学", "その座標の上で、時間・パラメータ変化をどう進めるか",
         ["Galerkin 射影（intrusive、物理）", "DMD（線形、データ）", "Operator Inference（多項式、データ）",
          "NN（任意、データ）"]),
        (12, "ROM", "＝ 基底 ＋ 力学", ["Acta Numerica 2026 の3要素：", "非線形パラメトリゼーション ＝ 基底",
                                      "縮約力学 ＝ 力学", "＋ オンラインソルバ"]),
    ]
    pw, gap = 9.4, 1.4
    for k, (idx, head, sub, items) in enumerate(content):
        ph = phs[idx]
        ph.left, ph.top = Cm(L + k * (pw + gap)), Cm(y + 0.2)
        ph.width, ph.height = Cm(pw), Cm(BODY_BOTTOM - y - 0.4)
        fill_tf(ph.text_frame, [{"t": head, "size": 32, "bold": True, "color": ACC1, "space_after": 4},
                                {"t": sub, "size": 18, "color": DEEP, "space_after": 14}]
                + [{"t": t, "bullet": not t.endswith("："), "color": DEEP} for t in items], size=20)
    for k, sym in enumerate(["＋", "＝"]):
        textbox(s, L + pw + k * (pw + gap), y + 0.2, gap, 1.6, [{"t": sym, "size": 32, "bold": True,
                                                                 "align": PP_ALIGN.CENTER}],
                color=ACC1, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "Hesthaven, Peherstorfer, Unger 2026")
    notes(s, "ROM は「どの座標で表すか（基底）」と「その座標でどう進めるか（力学）」の2つを選ぶこと。\n"
             "intrusive：ソルバの方程式・コードに手を入れる。autoencoder：入力を小さな潜在変数に圧縮して元に戻すよう学習する NN。\n"
             "【アニメーション】基底 → 力学 → ROM の順に")


def s16_galerkin(prs):
    s = new_slide(prs, LY_ONE, "Galerkin 射影", "非線形項の壁")
    y = claim(s, "支配方程式を POD モードに射影すると、物理に基づく r 次元の ODE（常微分方程式）になる。"
                 "ただし非線形項のコストは減らない")
    # a_dot (r x 1) = U_r^T (r x n) f( U_r (n x r) a (r x 1) )
    top, r, n = y + 0.6, 1.5, 4.8
    x = 3.0

    def mat(xx, w, h, name, col, lab):
        box(s, xx, top + (n - h) / 2, w, h, [{"t": name, "size": 16, "bold": True}], fill=col, radius_cm=0.04)
        textbox(s, xx - 0.6, top + (n + h) / 2 + 0.05, w + 1.2, 0.7, [{"t": lab, "align": PP_ALIGN.CENTER}],
                size=13, color=GREY)
        return xx + w

    def sym(xx, t, w=1.0):
        textbox(s, xx, top + n / 2 - 0.7, w, 1.4, [{"t": t, "size": 28, "align": PP_ALIGN.CENTER}],
                anchor=MSO_ANCHOR.MIDDLE)
        return xx + w

    x = mat(x, 0.55, r, "", ACC2, "ȧ：r×1") + 0.1
    textbox(s, 3.0 - 1.4, top + n / 2 - 0.5, 1.4, 1.0, [{"t": "ȧ", "size": 24, "bold": True,
                                                        "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    x = sym(x, "=")
    x = mat(x, n, r, "U_{r}^{T}", ACC1, "r×n") + 0.2
    fx = x
    x = sym(x, "f (", 1.4)
    x = mat(x, r, n, "U_{r}", ACC1, "n×r") + 0.75
    x = mat(x, 0.55, r, "", ACC2, "a：r×1") + 0.1
    x = sym(x, ")", 0.7)
    box(s, fx - 0.1, top - 0.35, x - fx + 0.2, n + 1.5, [], fill=None, line=WARM, dash=True, line_w=2)
    textbox(s, x + 0.4, top + 0.6, L + CW - x - 0.4, 3.6, [
        {"t": "ここは n 次元のまま", "bold": True, "color": RGBColor(0xC0, 0x7A, 0x1A), "size": 20},
        {"t": "U_{r}a で n 次元の場に戻し、全格子で f を計算してから縮める", "size": 16},
    ], size=16, anchor=MSO_ANCHOR.MIDDLE)
    ty = top + n + 1.4
    textbox(s, L, ty, CW, BODY_BOTTOM - ty, [
        {"t": "x ≈ U_{r} a(t) を支配方程式 ẋ = f(x) に代入し、U_{r}^{T} を掛ける → ȧ = U_{r}^{T} f(U_{r} a)",
         "bullet": True},
        {"t": "支配方程式の離散化オペレータに触る必要がある（**intrusive**：ソルバのコードに手を入れる）",
         "bullet": True},
        {"t": "非線形項 f(U_{r} a) は元の n 次元で評価する → DEIM の原論文の指摘："
              "POD は変数の数を減らすが、非線形項の評価コストは元の問題のまま", "bullet": True},
    ], size=20)
    cite(s, "B&K Ch.12｜Chaturantabut & Sorensen 2010")
    notes(s, "式の形を行列の大きさで見せる。U_r^T f(U_r a) の f の中は n 次元。ここが遅いままになる。\n"
             "次のスライドの DEIM が、この f を少数の点だけで評価する方法。\n"
             "【アニメーション】下の3項目：上から順に")


def s17_hyper(prs):
    one_gif_slide(
        prs, "Hyper-reduction", "Gappy POD と DEIM",
        "非線形項を「少数の点でだけ評価して補間する」ことで、コストを r に比例させる",
        ANIM / "17_deim_greedy.gif",
        [{"t": "Hyper-reduction：非線形項の計算そのものを減らす工夫", "size": 16, "color": GREY, "space_after": 10},
         {"t": "**Gappy POD**（Everson & Sirovich 1995）：欠けた場から、モード係数を最小二乗で復元。"
               "センサが少ない実験データの補完にも使える（演習②）", "bullet": True, "space_after": 8},
         {"t": "**DEIM**（Chaturantabut & Sorensen 2010）：非線形項のスナップショットに POD をかけ、"
               "**貪欲法**で補間点を1点ずつ選ぶ（残差が最大の点を次に選ぶ）", "bullet": True, "space_after": 8},
         {"t": "**QDEIM**：ピボット付き QR 分解で点を選ぶ（B&K §13.5–13.6）", "bullet": True}],
        "B&K Ch.13｜Everson & Sirovich 1995｜Chaturantabut & Sorensen 2010",
        "GIF 17：残差が最大の点を1つ選び、その点で合わせた後の残差でまた最大の点を選ぶ。点が後流に広がっていく。"
        "右は、選んだ点の数 j で非線形項を補間したときの誤差。\n"
        "貪欲法：その時点で一番よい選択を1つずつ積み重ねる方法。",
        gif_w=18.4, size=18)


def s18_intrusive(prs):
    s = new_slide(prs, LY_TWO, "Intrusive か Non-intrusive か", "商用ソルバでは、方程式に触れない手法を選ぶ")
    y = claim(s, "商用ソルバ（STAR-CCM+ 等）は方程式に触れない → 非侵入的（non-intrusive）手法を選ぶ")
    bh = BODY_BOTTOM - y - 1.9
    for x, col, head, sub, items in [
        (L, LIGHT, "Intrusive（侵入的）", "ソルバの方程式・コードに手を入れる",
         ["Galerkin 射影（＋ DEIM）", "自作コード・OpenFOAM 等なら可能"]),
        (17.54, HEAD_FILL, "Non-intrusive（非侵入的）", "ソルバの入出力データだけを使う",
         ["**Operator Inference**（Peherstorfer & Willcox 2016）：縮約力学を多項式と仮定し、"
          "係数行列を最小二乗で学習",
          "物理の形（二次項など）を保ち、ソルバのソースに触らない",
          "レビュー：Kramer et al. 2024（最大9桁の高速化を報告）",
          "その他：POD 係数の補間 / GP、NN（LSTM、neural ODE）"]),
    ]:
        box(s, x, y, 15.2, 1.6, [{"t": head, "size": 22, "bold": True, "space_after": 0},
                                 {"t": sub, "size": 15}], fill=col, line=ACC1, line_w=1.0, align=PP_ALIGN.LEFT, margin=0.35)
        textbox(s, x, y + 1.8, 15.2, bh - 0.1, [{"t": t, "bullet": True, "space_after": 6} for t in items], size=17)
    box(s, L, BODY_BOTTOM - 1.15, CW, 1.1, [{"t": "これが「実際にやる時に、最初に決めること」", "size": 20,
                                             "bold": True, "color": DEEP}], fill=WARM)
    cite(s, "Peherstorfer & Willcox 2016｜Kramer+ 2024")
    notes(s, "自分のソルバのソースに触れるかどうかで、選べる手法が半分に分かれる。STAR-CCM+ などの商用ソルバは触れない。\n"
             "Operator Inference：縮約した方程式が「一次項＋二次項」のような形だと仮定して、その係数だけをデータから求める。\n"
             "LSTM：時系列を扱う NN。neural ODE：微分方程式の右辺を NN で表すもの。\n"
             "【アニメーション】右の列：Intrusive の説明のあとで表示")


def s19_manifold(prs):
    s = new_slide(prs, LY_ONE, "線形部分空間の限界", "非線形多様体と、Deep Learning の存在意義")
    y = claim(s, "移流が支配的だと線形基底が効かない（Kolmogorov barrier）→ 非線形の基底（autoencoder）")
    gif = BUILD / "19_manifold_crop.gif"
    gh = BODY_BOTTOM - y - 0.4
    w, _ = picture(s, gif, L, y, h=gh)
    th = BODY_BOTTOM - y - 1.0
    tw = th * image_size(THESIS / "T7_rows3.png")[0] / image_size(THESIS / "T7_rows3.png")[1]
    picture(s, THESIS / "T7_rows3.png", L + CW - tw, y, h=th)
    textbox(s, L + CW - tw - 0.6, y + th + 0.05, tw + 0.6, 0.9, [
        {"t": "単螺旋 → 二重螺旋（Fig. 5.20）", "size": 11, "color": GREY, "align": PP_ALIGN.CENTER}], size=11)
    x = L + w + 0.4
    side_text(s, x, y, [
        {"t": "**Kolmogorov n-width**：最良の r 次元線形部分空間でも、誤差が遅くしか減らない問題がある",
         "bullet": True, "space_after": 8},
        {"t": "Lee & Carlberg (2020)：畳み込み autoencoder で非線形多様体を作り、その上で Galerkin / LSPG 射影",
         "bullet": True, "space_after": 8},
        {"t": "B&K §13.7 Decoder networks", "bullet": True, "space_after": 8},
        {"t": "Balanced POD（B&K Ch.9）：入出力（制御）が目的の ROM。今日は存在だけ", "bullet": True},
    ], size=16, w=L + CW - tw - 0.4 - x)
    cite(s, "Peherstorfer 2022｜Lee & Carlberg 2020｜Rowley & Dawson 2017｜B&K §13.7, Ch.9｜博論 Fig. 5.20")
    notes(s,
        "GIF 19：曲がった解の集まり（曲線）を、直線（線形部分空間）で近似しようとしても届かない。視点がゆっくり回る。\n"
        "LSPG：最小二乗 Petrov–Galerkin 射影。Balanced POD は入出力を残す制御向けの ROM で、今日は名前だけ。\n"
        "右の博論図：同じ運転条件でも渦ロープが単螺旋と二重螺旋を行き来する（図は t₁ から 2Δt まで。"
        "このあと二重螺旋から単螺旋に戻る）。両方を1つの線形基底で表すとモードが倍以上要る。")


def s20_map(prs):
    s = new_slide(prs, LY_TABLE, "手法マップ", "象限と制約から、手法が決まる", keep=(12,))
    y = claim(s, f"自分の問題（{pg('5')} の象限）と制約（intrusive にできるか）から、手法が決まる")
    ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
    rows = [
        ["基底 ＼ 力学", "Galerkin ＋ DEIM", "DMD", "OpInf（回帰）", "NN"],
        ["線形\n（POD）", "intrusive\nChaturantabut & Sorensen 2010", "non-intrusive\nSchmid 2010；Tu et al. 2014",
         "non-intrusive\nPeherstorfer & Willcox 2016", "non-intrusive\nPOD 係数を NN で進める\n（例：POD-LSTM、Zhang et al. 2026）"],
        ["非線形\n（AE・二次多様体）", "intrusive\nLee & Carlberg 2020\n（Galerkin / LSPG）", "—",
         "non-intrusive\n二次多様体：Geelen et al. 2023\nLaSDI 系：潜在 ODE を回帰\n（Bonneville et al. 2024）",
         "—"],
    ]
    table(s, L, y + 0.2, [5.6, 5.9, 5.2, 8.4, 6.5], rows, size=16, row_h=3.2, placeholder=ph,
          align_first=PP_ALIGN.CENTER, header_h=1.0)
    textbox(s, L, BODY_BOTTOM - 2.3, CW, 2.3, [
        {"t": f"使い方：① {pg('5')} で象限を決める（何が変わるか × 線形/非線形）→ ② 方程式に触れるかで列を選ぶ",
         "size": 18, "bullet": True},
        {"t": "基底 ＝ 非線形パラメトリゼーション、力学 ＝ 縮約力学（Hesthaven, Peherstorfer, Unger 2026 の3要素）",
         "size": 18, "bullet": True},
    ], size=18)
    cite(s, "各セルの文献")
    notes(s, "これまでの話を1枚にまとめる。AE：autoencoder。LaSDI：潜在空間の力学を同定する手法の系列。\n"
             "非線形基底 × DMD と 非線形基底 × NN の欄は、本講で扱う文献がないので空欄。\n"
             "LaSDI：AE で潜在空間を作り、その中の力学を多項式の ODE として回帰で同定する。"
             "力学の作り方が OpInf と同じ「回帰」なので、この列に置く。\n"
             "【アニメーション】表の行：線形（POD）→ 非線形（AE・二次多様体）の順に。最後に下の「使い方」")


def s21_timeline(prs):
    s = new_slide(prs, LY_FREE, "年表", "数学 → 流体 → 力学 → 非侵入 → 非線形")
    y = claim(s, "数学（SVD/PCA）→ 流体（POD）→ 力学（Galerkin, DMD）→ 非侵入（OpInf）→ "
                 "非線形（AE, 二次多様体）の順に進んできた")
    eras = {"数学": ACC5, "流体": ACC4, "力学": ACC1, "非侵入": ACC3, "非線形": WARM,
            "サーベイ": GREY}
    ev = [
        ("1901", "Pearson", "PCA", "数学"),
        ("1967", "Lumley", "POD（乱流）", "流体"),
        ("1987", "Sirovich", "method of snapshots", "流体"),
        ("1995", "Everson & Sirovich", "Gappy POD", "流体"),
        ("2010", "Chaturantabut & Sorensen", "DEIM", "力学"),
        ("2010", "Schmid", "DMD", "力学"),
        ("2014", "Tu et al.", "exact DMD", "力学"),
        ("2015", "Benner et al.", "サーベイ", "サーベイ"),
        ("2016", "Peherstorfer & Willcox", "OpInf", "非侵入"),
        ("2020", "Lee & Carlberg", "AE-ROM", "非線形"),
        ("2022", "Barnett & Farhat", "二次多様体", "非線形"),
        ("2022", "Hesthaven et al.", "Acta Numerica サーベイ", "サーベイ"),
        ("2024", "Kramer et al.", "OpInf レビュー", "非侵入"),
        ("2026", "Hesthaven, Peherstorfer, Unger", "非線形 ROM サーベイ", "サーベイ"),
    ]
    ax_y = y + 5.7
    x0, x1 = L + 1.2, L + CW - 1.7
    line(s, L, ax_y, L + CW, ax_y, color=GREY, width=2.5)
    step = (x1 - x0) / (len(ev) - 1)
    lw = 4.4
    for k, (yr, who, what, era) in enumerate(ev):
        cx = x0 + k * step
        col = eras[era]
        box(s, cx - 0.22, ax_y - 0.22, 0.44, 0.44, [], fill=col, shape=MSO_SHAPE.OVAL)
        up = k % 2 == 0
        ly = ax_y - 4.9 if up else ax_y + 0.6
        line(s, cx, ax_y - (0.25 if up else -0.25), cx, ly + (4.3 if up else 0.0), color=col, width=1.25)
        textbox(s, cx - lw / 2, ly, lw, 4.3, [
            {"t": yr, "size": 16, "bold": True, "color": col, "align": PP_ALIGN.CENTER, "space_after": 0},
            {"t": what, "size": 14, "bold": True, "align": PP_ALIGN.CENTER, "space_after": 0},
            {"t": who, "size": 12, "color": GREY, "align": PP_ALIGN.CENTER},
        ], size=14, anchor=MSO_ANCHOR.BOTTOM if up else MSO_ANCHOR.TOP, space_after=0)
    lx = L
    for name, col in eras.items():
        box(s, lx, BODY_BOTTOM - 0.55, 0.45, 0.45, [], fill=col, shape=MSO_SHAPE.OVAL)
        textbox(s, lx + 0.55, BODY_BOTTOM - 0.7, 3.0, 0.75, [{"t": name, "size": 14}], size=14,
                anchor=MSO_ANCHOR.MIDDLE)
        lx += 3.6
    textbox(s, lx + 0.5, BODY_BOTTOM - 0.7, L + CW - lx - 0.5, 0.75,
            [{"t": "横軸は等間隔（年の長さに比例しない）", "size": 13, "color": GREY,
              "align": PP_ALIGN.RIGHT}],
            size=13, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "Pearson 1901・Lumley 1967 は定説として扱う")
    notes(s, "左から右へ、数学 → 流体 → 力学 → 非侵入 → 非線形。色が分野。\n"
             "【アニメーション】年表の項目：左から順に1つずつ（1901 → 2026）")


def s22_frontier(prs):
    s = new_slide(prs, LY_THREE, "最前線", "3つの方向と、何が嬉しいか")
    y = claim(s, "最前線は「非線形基底」「非侵入的力学」「離散化非依存」の3方向に動いている")
    cw = (CW - 2 * GAP) / 3
    cols = [
        ("非線形基底", [
            "二次多様体（Barnett & Farhat 2022；Geelen, Wright, Willcox 2023）、"
            "NN 拡張射影（Barnett, Farhat, Maday 2023）",
            "→ **移流・衝撃波を少ないモードで表す。学習 CFD の本数とオンライン次元が減る**",
        ]),
        ("非侵入的力学", [
            "OpInf：ソルバのソースに触らず、物理の形を保った力学を学習 → **商用 CFD で使え、素の NN より外挿が安定**",
            "LaSDI 系（Bonneville et al. 2024）：潜在力学を同定してパラメータ補間、UQ 付き能動学習 → "
            "**高忠実度解析の本数を減らす**",
            "構造保存（Hamiltonian OpInf 等）→ **長時間回しても発散しない**",
        ]),
        ("離散化非依存", [
            "RONOM（arXiv 2507.12814）：ROM の誤差評価と neural operator の離散化非依存性を橋渡し → "
            "**PIV 格子と CFD 格子の混在、解像度をまたぐ**",
            "PDE 基盤モデル（PROSE-FD、BCAT 等）：事前学習 → 少データで微調整。**産業適用は検証段階**",
        ]),
    ]
    ch = BODY_BOTTOM - y - 1.4
    for k, (head, items) in enumerate(cols):
        x = L + k * (cw + GAP)
        box(s, x, y, cw, 1.3, [{"t": head, "size": 22, "bold": True}], fill=HEAD_FILL, line=ACC1, line_w=1.0,
            shape=MSO_SHAPE.PENTAGON)
        card(s, x, y + 1.5, cw, BODY_BOTTOM - 1.1 - (y + 1.5), None,
             [{"t": t, "bullet": True, "space_after": 8} for t in items], size=14)
    textbox(s, L, BODY_BOTTOM - 0.9, CW, 0.9, [
        {"t": "サーベイの系譜：Benner et al. 2015 → Hesthaven, Pagliantini, Rozza 2022 → "
              "Hesthaven, Peherstorfer, Unger 2026（後の2本は Acta Numerica）", "size": 13}], size=13,
            anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "Barnett & Farhat 2022｜Geelen+ 2023｜Barnett+ 2023｜Bonneville+ 2024｜RONOM", to="apxA2")
    notes(s, "各方向について「何が嬉しいか」（緑の太字）を中心に話す。\n"
             "neural operator：関数から関数への写像を学習する NN。格子の細かさに依らずに使える。UQ：不確かさの定量化。\n"
             "【アニメーション】3方向：左から順に")


def s23_workflow(prs):
    s = new_slide(prs, LY_AGENDA, "ワークフロー", "出力 → 精度 → 入力 → 手法 → ツール → 限界")
    y = claim(s, "設計は「出力 → 精度 → 入力 → 手法 → ツール → 限界」の順に決める")
    steps = [
        ("1  出力", "誰が、何の計算に、この値を使うか"),
        ("2  精度", "周波数が合えばいいのか、振幅も要るのか。レイテンシ（答えが出るまでの時間）は"),
        ("3  入力", "どのデータが何本あるか。実験か CFD か、センサはどこか"),
        ("4  手法", f"何が変わるか（{pg('5')} の象限）。方程式に触れるか（intrusive にできるか）"),
        ("5  ツール", "どのソフトで作り、誰がどこで動かすか"),
        ("6  限界", "どこで外れるか（外挿・分岐・ノイズ）を先に書く"),
    ]
    n = len(steps)
    rh = (BODY_BOTTOM - y - 0.2) / n
    for k, (head, q) in enumerate(steps):
        yy = y + k * rh
        box(s, L, yy, 5.4, rh - 0.35, [{"t": head, "size": 20, "bold": True}], fill=HEAD_FILL if k < 5 else WARM,
            line=ACC1 if k < 5 else None, line_w=1.0)
        if k < n - 1:
            arrow(s, L + 2.7, yy + rh - 0.35, L + 2.7, yy + rh, color=ACC1, width=2)
        textbox(s, L + 5.8, yy, 18.4, rh - 0.35, [{"t": f"「{q}」"}], size=20, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 26.0, y, L + CW - 26.0, 5.0, "使い方", [
        f"思考実験（{pg('24')}）で、この6項目を埋める",
        f"答え合わせ（{pg('25a')}）も同じ6項目の表。{pg('25b')}–{PAGE['25c']} は (A) の深掘り",
    ], size=15)
    notes(s, "ROM を作り始める前に、6つの問いに上から答える。手法は4番目。最初に手法を決めない。\n"
             "【アニメーション】6ステップ：上から順に")


def s24_task(prs):
    s = new_slide(prs, LY_FOUR, "思考実験：課題", "2つの例題から1つ選び、ROM の設計書を書く")
    y = claim(s, f"データ（①〜④）から出発して、例題 (A) か (B) の ROM 設計書（{pg('23')} の6項目）を書く")
    qw = 15.2
    rh = 0.9
    rows = [["データ ＼ 例題", "(A) 渦の脈動", "(B) 電池の熱"]] + [[d, "", ""] for d in
            ["① 壁面圧力（数点のセンサ）", "② PIV（2D 断面、時間分解）", "③ CFD の 2D 断面", "④ 3D の結果そのまま"]]
    table(s, L, y, [7.0, 4.1, 4.1], rows, size=14, row_h=rh)
    ey = y + 5 * rh + 0.4
    card(s, L, ey, qw, BODY_BOTTOM - ey, "例題：何を作るか", [
        {"t": "(A) 水力発電所：渦の脈動を系統モデルに入れる", "bold": True, "space_after": 2},
        {"t": "水車の出口の管（ドラフトチューブ）で渦が振れ回り、壁の圧力が脈動する。"
              "この脈動を「揺れの源」として、発電所全体の 1D 水撃モデルに入れる", "size": 14, "space_after": 2},
        {"t": f"1D 水撃モデル：管の中を圧力波が伝わる様子を、管の長さ方向だけで解くモデル（{pg('3')} 右の 1D 系統モデルの一種）",
         "size": 13, "color": GREY, "space_after": 10},
        {"t": "(B) バッテリーパックの熱：3D 解析を 1D の部品にする", "bold": True, "space_after": 2},
        {"t": f"急速充電中のセル温度を、系統モデルの中で速く計算する（{pg('3')} の real-time の例）", "size": 14},
    ], size=14)
    rx = 17.54
    textbox(s, rx, y, qw, 0.8, [{"t": "(A) か (B) を選び、10分で6項目を埋める → 2組が発表", "bold": True,
                                 "color": ACC1}], size=15)
    # T6 and T5b share one row; T5a_2 spans the row below. Scale both rows to fit the column.
    fy = y + 1.0
    fh = BODY_BOTTOM - fy - 0.7
    a6, a5b, a5a = (image_size(THESIS / f)[0] / image_size(THESIS / f)[1] for f in ("T6.png", "T5b.png", "T5a_2.png"))
    h1 = (qw - 0.3) / (a6 + a5b)
    h2 = qw / a5a
    k = min(1.0, (fh - 0.2) / (h1 + h2))
    h1, h2 = h1 * k, h2 * k
    x0 = rx + (qw - (h1 * (a6 + a5b) + 0.3)) / 2
    w1, _ = picture(s, THESIS / "T6.png", x0, fy, h=h1)
    picture(s, THESIS / "T5b.png", x0 + w1 + 0.3, fy, h=h1)
    picture(s, THESIS / "T5a_2.png", rx + (qw - h2 * a5a) / 2, fy + h1 + 0.2, h=h2)
    textbox(s, rx, BODY_BOTTOM - 0.65, qw, 0.65, [
        {"t": "(A) の題材：渦糸（Fig. 3.1）・PIV 配置（Fig. 3.8）・圧力センサ（Fig. 3.3）",
         "size": 12, "color": GREY}], size=12, anchor=MSO_ANCHOR.MIDDLE)
    cite(s, "博論 Fig. 3.1, 3.3, 3.8", to="apxA3")
    notes(s, f"進め方：グループで (A) か (B) を選び、出発するデータ（①〜④）を決めて、{pg('23')} の6項目"
             "（出力・精度・入力・手法・ツール・限界）を10分で埋める。\n"
             "(A) は博論の題材。右下の図で、渦の振れ回り・PIV・圧力センサの配置を見せる。\n"
             f"水撃：管の中の流れが急に変わると、圧力波が管を往復する現象。1D 水撃モデルはこれを管の長さ方向だけで解く。"
             f"{pg('3')} 右の「ポンプ・管路・水車・制御」をつないだ系統モデルの管路の部分にあたる。\n"
             "① の壁面圧力は、振動騒音の対象になる量。\n"
             "【アニメーション】マトリクス → 例題（(A)・(B)）→ (A) の図 の順に")


def s25a_answers(prs):
    rows = [
        ["", "(A) 渦の脈動 → 1D 水撃モデル", "(B) バッテリー熱 3D → 1D"],
        ["出力", "励振源（運動量・質量源）とドラフトチューブのパラメータ（波速など）→ 系統の共振解析",
         "セル温度 → SIL/HIL・系統モデル（Siemens）／各セルの平均温度（Ansys）"],
        ["精度・\nレイテンシ", "固有振動数と共振条件が合うこと。レイテンシは問わない（設計段階）",
         "3.92 h → 0.04 s（急速充電1サイクル、Siemens）／誤差約 1.4% で 10,000 倍（LTI、Ansys）"],
        ["入力データ", "縮小模型の壁面圧力 ＋ 2D-PIV、LDV、外部加振。流量と Thoma 数を掃引",
         "3D 解析6本（流量 0.025–0.1 kg/s、入口 15–21 °C）／CFD のステップ応答"],
        ["手法", "1D 分布定数モデル ＋ パラメータ同定。励振源の強さを渦の軌道半径・周波数・循環に結ぶ",
         "NN（温度）＋ 2-RC 等価回路／LTI 状態空間 ＋ 流量で LPV 補間"],
        ["ツール", "SIMSEN、自前の PIV/LDV 処理", "STAR-CCM+、Simcenter ROM Builder、Amesim／Fluent・Icepak、Twin Builder"],
        ["限界", "外挿：ある流量以下で渦が急に構造を失い、強度が激減（分岐）",
         "外挿：流量を変えると LTI の仮定が破れる"],
    ]
    s = new_slide(prs, LY_TABLE, "思考実験：答え合わせ (1/3)",
                  "答えの型は「物理の骨組み ＋ 少数パラメータの同定」", keep=(12,))
    y = claim(s, "答えの型は「物理の骨組み＋少数パラメータを実験/CFDから同定」。"
                 "純粋なブラックボックスは診断用か、固定運転点での再構成に留まる")
    ph = [q for q in s.placeholders if q.placeholder_format.idx == 12][0]
    table(s, L, y + 0.1, [3.6, 14.0, 14.0], rows, size=16, header_size=17, row_h=1.55, placeholder=ph,
          align_first=PP_ALIGN.CENTER, header_h=0.9)
    cite(s, "(A) Favrel+ 2015, Landry+ 2016, Alligné+ 2016｜(B) Siemens 2023, Hu+（Ansys）", to="apxA3")
    notes(s, "各グループの答えと、この表の行を1つずつ比べる。\n"
             "(B) の数値はベンダーの自己申告。\n"
             f"(A) の物理の骨組みは次の {pg('25b')}、ブラックボックスとの比較は {pg('25c')}。\n"
             "【アニメーション】列 (A) → (B) の順に")


def s25b_skeleton(prs):
    s = new_slide(prs, LY_TABLE, "思考実験：答え合わせ (2/3)", "(A) の物理の骨組み：2つのパラメータで脈動が決まる")
    y = claim(s, "軌道半径と循環の2パラメータで、壁面脈動の振幅が決まる → これが 1D モデルに渡す励振源")
    w1, h1 = picture(s, THESIS / "T4a.png", L, y, w=15.2)
    textbox(s, L, y + h1 + 0.05, 15.2, 1.2, [
        {"t": "偏心した渦と壁面センサの幾何、圧力の波形（Fig. 3.31）", "size": 14, "color": GREY}], size=14)
    rx = 17.54
    w2, h2 = picture(s, THESIS / "T4b.png", rx, y, w=15.2)
    textbox(s, rx, y + h2 + 0.05, 15.2, 1.8, [
        {"t": "理論曲線 C_{p}^{*}(r^{*}) と全実験点（Fig. 3.33）。実験範囲 r^{*} < 0.3 では **C_{p}^{*} ≈ r^{*}**（式 3-40）",
         "size": 15},
        {"t": "r^{*}：渦の軌道半径（管の半径で無次元化）。理論は r^{*} = 1 で発散し、壁の近くでは近似の限界（博論 p.76）",
         "size": 13, "color": GREY},
    ], size=15)
    cy = y + h2 + 2.9
    box(s, rx, cy, 15.2, min(3.2, BODY_BOTTOM - cy), [
        {"t": "1D 水撃モデルに渡すもの", "bold": True, "color": ACC1, "size": 16, "space_after": 6},
        {"t": "PIV で渦の軌道半径 r^{*} と循環を測る → 壁面脈動の振幅（C_{p}^{*} ≈ r^{*}）→ 励振源の強さ", "size": 15},
    ], fill=HEAD_FILL, line=ACC1, line_w=1.0, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.35)
    cite(s, "博論 Fig. 3.31, 3.33, p.74–77", to="apxA3")
    notes(s, "(A) の答えの核：物理モデル（偏心した渦の圧力場）が骨組み。決めるパラメータは軌道半径と循環の2つ。"
             "これを PIV から同定すれば、1D モデルに渡す励振源になる。EPFL の Favrel らが励振源の強さを軌道半径・周波数・循環に結び付けたのと同じ考え方。\n"
             f"次の {pg('25c')} で、ブラックボックス（POD-LSTM）と比べる。")


def s25c_blackbox(prs):
    s = new_slide(prs, LY_TABLE, "思考実験：答え合わせ (3/3)", "比較：ブラックボックス（POD-LSTM、Zhang et al. 2026）")
    y = claim(s, "壁のセンサだけで学習したブラックボックスは、センサから見えないモードを予測できない")
    fh = BODY_BOTTOM - y - 1.0
    fw = fh * image_size(PAPERS / "zhang2026_fig17.png")[0] / image_size(PAPERS / "zhang2026_fig17.png")[1]
    fx = L + CW - fw
    picture(s, PAPERS / "zhang2026_fig17.png", fx, y, h=fh)
    textbox(s, fx - 0.3, y + fh + 0.02, fw + 0.3, 0.95, [
        {"t": "Zhang et al. 2026, Fig. 17：PL 条件の POD 係数。青＝真値、赤破線＝LSTM の予測。"
              "モード3（R² = −3.366）は外れる", "size": 12, "color": GREY}], size=12)
    lw = fx - L - 0.5
    rh = 4.0
    rw, _ = picture(s, BUILD / "zhang2026_fig2_ring.png", L, y, h=rh)
    textbox(s, L + rw + 0.3, y, lw - rw - 0.3, rh, [
        {"t": "LES（3D 非定常 CFD）で3つの運転点（PL・BEP・HL）を計算", "bullet": True, "space_after": 3},
        {"t": "壁面圧力センサ4点（左図の赤）× 20 時刻 → LSTM → POD 係数 → 3D 圧力場を再構成", "bullet": True,
         "space_after": 3},
        {"t": "運転点ごとに、同じ計算の中で学習と評価（180 組を 8:2 に分割）", "bullet": True},
    ], size=14, anchor=MSO_ANCHOR.MIDDLE)
    qy = y + rh + 0.3
    qh = 3.9
    box(s, L, qy, lw, qh, [
        {"t": "要旨から（原文）", "bold": True, "color": ACC1, "size": 13, "space_after": 4},
        {"t": "“PL Mode 3 and HL Mode 2, despite their regular temporal oscillations, yield strongly negative R² "
              "values (−3.366 and −186.6) because their spatial structures are concentrated away from the wall.”",
         "size": 13, "space_after": 6},
        {"t": "→ 規則正しいモードでも、壁から離れた構造はセンサに映らず、予測に失敗", "size": 13},
    ], fill=CARD, line=GREY, line_w=0.75, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.3)
    cy = qy + qh + 0.3
    card(s, L, cy, lw, BODY_BOTTOM - cy, "物理の骨組み（" + pg("25b") + "）との違い", [
        {"t": "再構成の L2 誤差：17.01%（PL）・7.17%（BEP）・12.91%（HL）", "bullet": True, "space_after": 3},
        {"t": "学習できるのはセンサに映るモードだけ。運転点をまたぐ予測は扱っていない", "bullet": True,
         "space_after": 3},
        {"t": "骨組みは、測る量（軌道半径・循環）を物理で先に選ぶ", "bullet": True},
    ], size=14)
    cite(s, "Zhang+ 2026, Energies 19(10):2300, Fig. 2・17・要旨（CC BY）", to="apxA3")
    notes(s, "比べる相手：Zhang et al. (2026)。フランシス水車のドラフトチューブを LES で計算し、壁の圧力センサ4点から"
             "POD 係数を LSTM で予測して、3D の圧力場を再構成した。\n"
             "PL：部分負荷、BEP：最高効率点、HL：高負荷。R²：決定係数。1 が完全一致、負は平均値で当てるより悪い。\n"
             "右の図：モード1・2 は真値に沿うが、モード3 は真値が右上がりなのに予測はほぼ一定。"
             "論文はこれを、モード3 の構造が壁から離れていてセンサに映らないためと説明している。\n"
             "図と引用文は CC BY ライセンス（出典を示せば転載可）。")


def s26_summary(prs):
    s = new_slide(prs, LY_STATEMENT, None)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Cm(L), Cm(1.6), Cm(CW), Cm(3.4)
    tf = t.text_frame
    add_runs(tf.paragraphs[0], "まずは手元のデータで POD/DMD を回し、自分の問題の象限を確認するところから", 30, ACC1,
             bold=True)
    y = 5.6
    cw = (CW - 2 * GAP) / 3
    card(s, L, y, cw, BODY_BOTTOM - y, "実装の道具（一覧）", [
        {"t": t, "bullet": True} for t in [
            "numpy の SVD／scipy", "PyDMD", "sklearn の PCA",
            "Operator Inference（Willcox グループの公開コード）", "PyTorch の autoencoder",
            "Simcenter ROM Builder・Ansys Twin Builder（商用）"]
    ], size=15)
    card(s, L + cw + GAP, y, cw, BODY_BOTTOM - y, "今日の4ゴールへの答え", [
        {"t": t, "bullet": True, "space_after": 6} for t in [
            "① いつ考えるか：同じ解析を何度も・実時間で回したく、データが低ランクのとき",
            "② 長所と短所：桁違いに速い／学習データが要り、外挿に弱い",
            "③ どこまで期待できるか：σ の減衰と象限（何が変わるか × 線形/非線形）で見積もる",
            "④ 何を準備するか：出力 → 精度 → 入力 → 手法 → ツール → 限界 の6項目"]
    ], size=14)
    card(s, L + 2 * (cw + GAP), y, cw, BODY_BOTTOM - y, "研究で使ってみよう", [
        {"t": "自分の研究のデータで、まず SVD を1回回してみてください。σ の並びを見るだけで、"
              "ROM が効きそうか分かります", "space_after": 14},
        {"t": "次回予告", "bold": True, "color": ACC1, "size": 15},
        "次回（シリーズの2）：Geometric Deep Learning（グラフやメッシュの上で学習する深層学習）",
    ], size=15)
    notes(s, "4つのゴールに1つずつ答えを返して締める。道具は名前の羅列だけ（注意点は来年度）。\n"
             "【アニメーション】3枚のカード：左から順に")


def appendix_refs(prs):
    refs = [
        "Brunton & Kutz, Data-Driven Science and Engineering, 2nd ed., Cambridge UP, 2022",
        "Benner, Gugercin, Willcox, \"A survey of projection-based model reduction methods for parametric "
        "dynamical systems\", SIAM Review 57(4):483–531, 2015",
        "Hesthaven, Pagliantini, Rozza, \"Reduced basis methods for time-dependent problems\", "
        "Acta Numerica 31:265–345, 2022",
        "Hesthaven, Peherstorfer, Unger, \"Nonlinear model reduction for transport-dominated problems\", "
        "Acta Numerica, 2026, DOI 10.1017/s0962492926100294（arXiv 2602.01397）",
        "Peherstorfer, \"Breaking the Kolmogorov Barrier with Nonlinear Model Reduction\", Notices AMS 69(5):725–, 2022",
        "Kramer, Peherstorfer, Willcox, \"Learning nonlinear reduced models from data with operator inference\", "
        "Annu. Rev. Fluid Mech. 56:521–548, 2024",
        "Peherstorfer & Willcox, \"Data-driven operator inference for nonintrusive projection-based model "
        "reduction\", CMAME 306:196–215, 2016",
        "Sirovich, \"Turbulence and the dynamics of coherent structures. I. Coherent structures\", "
        "Quart. Appl. Math. 45(3):561–571, 1987",
        "Everson & Sirovich, \"Karhunen-Loeve procedure for gappy data\", JOSA A 12(8):1657–1664, 1995",
        "Chaturantabut & Sorensen, \"Nonlinear Model Reduction via Discrete Empirical Interpolation\", "
        "SIAM J. Sci. Comput. 32(5):2737–2764, 2010",
        "Schmid, \"Dynamic mode decomposition of numerical and experimental data\", JFM 656:5–28, 2010",
        "Tu, Rowley, Luchtenburg, Brunton, Kutz, \"On dynamic mode decomposition: Theory and applications\", "
        "J. Comput. Dyn. 1(2):391–421, 2014",
        "Kutz, Brunton, Brunton, Proctor, Dynamic Mode Decomposition: Data-Driven Modeling of Complex Systems, "
        "SIAM, 2016",
        "Schmid, \"Dynamic mode decomposition and its variants\", Annu. Rev. Fluid Mech. 54:225–254, 2022",
        "Dawson, Hemati, Williams, Rowley, \"Characterizing and correcting for the effect of sensor noise in the "
        "dynamic mode decomposition\", Exp. Fluids 57:42, 2016",
        "Rowley & Dawson, \"Model Reduction for Flow Analysis and Control\", Annu. Rev. Fluid Mech. 49:387–417, 2017",
        "Taira et al., \"Modal Analysis of Fluid Flows: An Overview\", AIAA J 55(12), 2017, DOI 10.2514/1.J056060",
        "Taira et al., \"Modal analysis of fluid flows: Applications and outlook\", AIAA J 58(3):998–1022, 2020",
        "Lee & Carlberg, \"Model reduction of dynamical systems on nonlinear manifolds using deep convolutional "
        "autoencoders\", JCP 404:108973, 2020",
        "Barnett & Farhat, \"Quadratic approximation manifold for mitigating the Kolmogorov barrier in nonlinear "
        "projection-based model order reduction\", JCP 464:111348, 2022",
        "Geelen, Wright, Willcox, \"Operator inference for non-intrusive model reduction with quadratic "
        "manifolds\", CMAME 403:115717, 2023",
        "Barnett, Farhat, Maday, \"Neural-network-augmented projection-based model order reduction for "
        "mitigating the Kolmogorov barrier to reducibility\", JCP 492:112420, 2023",
        "Benner et al. (eds.), Model Order Reduction, 3 vols., De Gruyter, 2021（Open Access）",
        "Bellman, Dynamic Programming, Princeton UP, 1957",
        "Bonneville et al., \"A comprehensive review of latent space dynamics identification algorithms…\", "
        "arXiv 2403.10748, 2024",
        "RONOM: \"Reduced-Order Neural Operator Modeling\", arXiv 2507.12814（v2, 2026年1月）",
    ]
    for part, (a, b) in enumerate([(0, 13), (13, 26)], 1):
        s = new_slide(prs, LY_TABLE, "付録A　参考文献", f"({part}/3) 本編の文献", keep=(12,))
        ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
        rows = [["#", "文献"]] + [[str(i + 1), refs[i]] for i in range(a, b)]
        table(s, L, BODY_TOP, [1.2, CW - 1.2], rows, size=12, row_h=0.55, placeholder=ph,
              align_first=PP_ALIGN.CENTER)
    s = new_slide(prs, LY_TABLE, "付録A　参考文献", "(3/3) 思考実験の答え合わせ", keep=(12,))
    ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
    ex = [
        ("A", "Favrel et al., Exp. Fluids 56:215, 2015"),
        ("A", "Landry et al., J. Hydraulic Res. 54:185–196, 2016"),
        ("A", "Alligné, Nicolet, Béguin et al., \"Hydroelectric System Response to Part Load Vortex Rope "
              "Excitation\", 28th IAHR Symposium on Hydraulic Machinery and Systems, 2016"),
        ("A", "Zhang, Ma, Li, Kong, Xu, Huang, Wang, \"Reconstruction and Prediction of Three-Dimensional Transient "
              "Flow Field in a Draft Tube of Francis Turbine Using Sparse Sensors and a POD-LSTM Network\", "
              "Energies 19(10):2300, 2026, DOI 10.3390/en19102300（CC BY）"),
        ("B", "N. Ziegler, Siemens Simcenter blog, 15 Feb 2023（Simcenter ROM Builder によるバッテリー熱 ROM）"),
        ("B", "X. Hu et al., Ansys（LTI / LPV ROM、IMECE2011-62022）；X. Hu & Ma, Ansys blog, 2026"),
    ]
    rows = [["例題", "文献"]] + [[a, b] for a, b in ex]
    table(s, L, BODY_TOP, [2.0, CW - 2.0], rows, size=13, row_h=0.7, placeholder=ph, align_first=PP_ALIGN.CENTER)
    textbox(s, L, 13.6, CW, 2.4, [
        {"t": "劉志豪, 早稲田大学博士論文「管路内旋回流れによる渦心振れ回り現象に関する研究」, 2022年2月"
              "（スライドでは「劉 (2022) 博士論文」「博論」と表記。頁は製本の印刷ページ）", "size": 13},
        {"t": f"前回講義：早稲田大学「数値計算の“その先”」, 2026-07-17（{pg('4')} で言及）", "size": 13},
    ], size=13)


def appendix_chapters(prs):
    s = new_slide(prs, LY_TABLE, "付録B　教科書の章対応", "教科書のコードは1版向け。ROM の章は番号が1つずれる",
                  keep=(12,))
    ph = [p for p in s.placeholders if p.placeholder_format.idx == 12][0]
    rows = [
        ["内容（2版の章題）", "1版（2019）", "2版（2022）", "今日のページ"],
        ["Singular Value Decomposition (SVD)", "Ch.1", "Ch.1", f"{PAGE['4']}, {PAGE['6']}〜{PAGE['10']}"],
        ["Data-Driven Dynamical Systems（DMD）", "Ch.7", "Ch.7", f"{PAGE['11']}, {PAGE['12']}"],
        ["Balanced Models for Control", "Ch.9", "Ch.9", f"{PAGE['19']}"],
        ["Reduced-Order Models (ROMs)", "Ch.11", "Ch.12", f"{PAGE['16']}"],
        ["Interpolation for Parametric Reduced-Order Models（DEIM など）", "Ch.12", "Ch.13",
         f"{PAGE['17']}, {PAGE['19']}"],
    ]
    table(s, L, BODY_TOP, [15.6, 4.5, 4.5, 7.0], rows, size=18, row_h=1.1, placeholder=ph)
    textbox(s, L, 11.3, CW, 5.8, [
        {"t": "教科書のコード：github.com/dynamicslab/databook_python（2019年の1版向け）", "bullet": True},
        {"t": "講義動画・データ：databookuw.com", "bullet": True},
        {"t": "2版の公式 PDF：faculty.washington.edu/sbrunton/DataBookV2.pdf", "bullet": True},
        {"t": "演習のリポジトリ：github.com/ShigoZhihao/waseda-course-rom-1", "bullet": True},
    ], size=18)


# ---------------------------------------------------------------- whole deck
def set_east_asian_font(prs, typeface=EA_FONT):
    """Add <a:ea typeface=...> to every run property in the slides (spec 8.1, item 5)."""
    after = {qn(t) for t in ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
    for slide in prs.slides:
        for tag in ("a:r", "a:fld"):
            for r in slide._element.iter(qn(tag)):
                if r.find(qn("a:rPr")) is None:
                    r.insert(0, r.makeelement(qn("a:rPr"), {}))
        for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
            for rpr in slide._element.iter(qn(tag)):
                ea = rpr.find(qn("a:ea"))
                if ea is None:
                    ea = rpr.makeelement(qn("a:ea"), {})
                    nxt = next((c for c in rpr if c.tag in after), None)
                    if nxt is None:
                        rpr.append(ea)
                    else:
                        nxt.addprevious(ea)
                ea.set("typeface", typeface)


def set_japanese_line_breaks(prs):
    """Mark runs as Japanese and turn on the Japanese line-break rules.

    The template's default language is en-US. With it, PowerPoint can start a line with 。 or ）.
    """
    for slide in prs.slides:
        for p in slide._element.iter(qn("a:p")):
            pPr = p.find(qn("a:pPr"))
            if pPr is None:
                pPr = p.makeelement(qn("a:pPr"), {})
                p.insert(0, pPr)
            pPr.set("eaLnBrk", "1")
            pPr.set("hangingPunct", "1")
        for tag in ("a:rPr", "a:endParaRPr"):
            for rpr in slide._element.iter(qn(tag)):
                rpr.set("lang", "ja-JP")
                rpr.set("altLang", "en-US")


def clean_package(prs):
    """Remove what the template carries from its source deck: comment authors and the old slide titles."""
    drop = ("/relationships/commentAuthors", "/relationships/authors")
    for rid, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith(drop):
            prs.part.drop_rel(rid)
    # custom properties and the classification label (sensitivity label) of the template
    pkg = prs.part.package
    for rid, rel in list(pkg._rels.items()):
        if rel.reltype.endswith(("/custom-properties", "/classificationlabels")):
            pkg.drop_rel(rid)
    n_notes = sum(1 for sl in prs.slides if sl.has_notes_slide)
    app = next(pt for pt in prs.part.package.iter_parts() if pt.partname == "/docProps/app.xml")
    app._blob = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
        '<PresentationFormat>Widescreen</PresentationFormat>'
        f'<Slides>{len(prs.slides)}</Slides><Notes>{n_notes}</Notes><HiddenSlides>0</HiddenSlides>'
        '<MMClips>0</MMClips><ScaleCrop>false</ScaleCrop><LinksUpToDate>false</LinksUpToDate>'
        '<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>'
        '<AppVersion>16.0000</AppVersion></Properties>'
    ).encode("utf8")


def build():
    prepare_images()
    prs = Presentation(TEMPLATE)
    s01_title(prs)
    chapter(prs, "A", "動機", "なぜ ROM か、「削減」とは何か、何が変わるか（10分）")
    s02_goals(prs)
    s03_why(prs)
    s04_lowrank(prs)
    s05_quadrant(prs)
    chapter(prs, "B", "SVD → POD / PCA", "行列を3つに分けて、流れの形を取り出す（15分）")
    s06_svd(prs)
    s07_meaning(prs)
    s08_why(prs)
    s09_pod_pca(prs)
    s10_truncation(prs)
    chapter(prs, "C", "DMD", "データから線形の力学を取り出す（8分）")
    s11_dmd1(prs)
    s12_dmd2(prs)
    chapter(prs, "D", "演習", "Colab で、円柱後流と PIV 風データを解く（15分）")
    s13_ex1(prs)
    s14_ex2(prs)
    s14b_thesis(prs)
    chapter(prs, "E", "手法マップ", "ROM ＝ 基底 ＋ 力学（12分）")
    s15_basis_dynamics(prs)
    s16_galerkin(prs)
    s17_hyper(prs)
    s18_intrusive(prs)
    s19_manifold(prs)
    s20_map(prs)
    chapter(prs, "F", "歴史と最前線", "どこから来て、どこへ向かうか（7分）")
    s21_timeline(prs)
    s22_frontier(prs)
    chapter(prs, "G", "ワークフローと思考実験", "ROM の設計書を書いてみる（25分）")
    s23_workflow(prs)
    s24_task(prs)
    s25a_answers(prs)
    s25b_skeleton(prs)
    s25c_blackbox(prs)
    chapter(prs, "H", "まとめ", "実装の道具・今日のゴール・次回予告・質疑（8分）")
    s26_summary(prs)
    appendix_refs(prs)
    appendix_chapters(prs)
    assert len(prs.slides) == len(ORDER), (len(prs.slides), len(ORDER))
    for shape, key in LINKS:
        shape.click_action.target_slide = prs.slides[PAGE[key] - 1]
    set_east_asian_font(prs)
    set_japanese_line_breaks(prs)
    cp = prs.core_properties
    cp.author = cp.last_modified_by = "Zhihao Liu"
    cp.title = "Reduced Order Modeling（ROM）：流れ場の低次元化と予測"
    cp.subject = "流体機械演習B 第4回（3回入門シリーズの1）"
    cp.keywords = cp.comments = cp.category = ""
    clean_package(prs)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT.relative_to(ROOT)} ({len(prs.slides)} slides)")


def export_pdf(pptx, pdf):
    """Export with PowerPoint (COM). PowerPoint stays open if the user has other files open."""
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(pptx), ReadOnly=True, WithWindow=False)
    try:
        pres.SaveAs(str(pdf), 32)  # 32 = ppSaveAsPDF
    finally:
        pres.Close()
        if app.Presentations.Count == 0:
            app.Quit()
    print(f"saved {pdf.relative_to(ROOT)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pdf", action="store_true", help="also export the PDF with PowerPoint")
    a = ap.parse_args()
    build()
    if a.pdf:
        export_pdf(OUT, OUT.with_suffix(".pdf"))


if __name__ == "__main__":
    main()
